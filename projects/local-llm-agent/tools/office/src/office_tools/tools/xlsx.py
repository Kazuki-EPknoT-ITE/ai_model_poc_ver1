"""xlsx.*（protocol §3.3）。"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.chart import BarChart, LineChart, PieChart, Reference, ScatterChart, Series
from openpyxl.styles import Font
from openpyxl.utils.cell import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from office_tools.envelope import Item, Outcome, paginate
from office_tools.errors import invalid
from office_tools.fsutil import (
    atomic_write,
    check_expect_hash,
    ensure_absent,
    existing_file,
    ext,
    resolve,
    sha256_file,
)
from office_tools.registry import ToolContext, tool
from office_tools.scalars import scalar_to_python, to_json
from office_tools.xlsx import recalc as rc
from office_tools.xlsx.book import check_ext, extent, has_formulas, mark_full_calc, open_book, sheet
from office_tools.xlsx.refs import CellRange, col_index, parse_cell, parse_range, quote_sheet, split_sheet
from office_tools.xlsx.shift import Shift, apply_shift, rename_sheet_refs

_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


# ---------------------------------------------------------------- 共通


def _save(
    path: str, real: Path, wb: Workbook, do_recalc: bool
) -> tuple[dict[str, Any], dict[str, Any] | None, list[str]]:
    """一時ファイルに保存（必要なら再計算）してから置き換える。"""
    result: dict[str, Any] | None = None
    warnings: list[str] = []
    formulas = has_formulas(wb)
    if formulas and (not do_recalc or not rc.engine_available()):
        mark_full_calc(wb)

    def writer(tmp: Path) -> None:
        nonlocal result
        wb.save(tmp)
        if do_recalc and formulas:
            result = rc.recalc_path(tmp)
            if result is None:
                warnings.append(rc.NO_ENGINE_WARNING)

    artifact = atomic_write(path, real, writer)
    return artifact, result, warnings


def _check_sheet_name(name: str) -> None:
    if _BAD_SHEET_CHARS.search(name) or name.startswith("'") or name.endswith("'"):
        raise invalid(f"シート名 {name} に使えない文字が含まれます（[ ] : * ? / \\）")


def _write_value(ws: Worksheet, col: int, row: int, value: Any, number_format: str | None = None) -> None:
    cell = ws.cell(row=row, column=col)
    if isinstance(cell, MergedCell):
        raise invalid(f"{cell.coordinate} は結合セルの一部です。結合範囲の左上のセルに書いてください")
    py = scalar_to_python(value)
    cell.value = py
    if isinstance(py, str) and py.startswith("="):
        cell.data_type = "s"  # 文字列として書く（数式にしない）
    if number_format:
        cell.number_format = number_format
    elif isinstance(value, dict) and cell.number_format == "General":
        cell.number_format = "yyyy-mm-dd"


def _load_for_write(args: dict[str, Any]) -> tuple[str, Path, Workbook]:
    path = args["path"]
    real = existing_file(path)
    check_ext(path, real)
    check_expect_hash(real, args.get("expect_hash"))
    return path, real, open_book(path, real)


def _write_result(artifact: dict[str, Any], recalc: dict[str, Any] | None, **data: Any) -> dict[str, Any]:
    out: dict[str, Any] = {"artifact": artifact, **data}
    if recalc is not None:
        out["recalc"] = recalc
    return out


# ---------------------------------------------------------------- 読み取り


def _zip_flags(real: Path) -> bool:
    with zipfile.ZipFile(real) as z:
        return any(n.lower().endswith("vbaproject.bin") for n in z.namelist())


def _pivots(ws: Any) -> list[dict[str, Any]]:
    """ピボットテーブルの定義（XLS-09: 読み取りのみ）。集計値はシート上の location を xlsx.read_range で読む。"""
    out: list[dict[str, Any]] = []
    for pt in getattr(ws, "_pivots", None) or []:
        cache = getattr(pt, "cache", None)
        names = [f.name for f in (getattr(cache, "cacheFields", None) or [])]

        def field(i: Any, names: list[str] = names) -> str:
            return names[i] if isinstance(i, int) and 0 <= i < len(names) else f"#{i}"

        src = getattr(getattr(cache, "cacheSource", None), "worksheetSource", None)
        out.append(
            {
                "name": pt.name,
                "location": pt.location.ref if pt.location is not None else None,
                "source": {"sheet": src.sheet, "ref": src.ref, "name": src.name} if src is not None else None,
                "rows": [field(f.x) for f in pt.rowFields if f.x is not None and f.x >= 0],
                "columns": [field(f.x) for f in pt.colFields if f.x is not None and f.x >= 0],
                "filters": [field(f.fld) for f in pt.pageFields],
                "values": [
                    {"name": d.name, "field": field(d.fld), "function": d.subtotal or "sum"}
                    for d in pt.dataFields
                ],
            }
        )
    return out


@tool("xlsx.info")
def xlsx_info(ctx: ToolContext, args: dict[str, Any]) -> Outcome:
    path = args["path"]
    real = existing_file(path)
    wb = open_book(path, real)
    sheets = []
    for ws in wb.worksheets:
        sheets.append(
            {
                "name": ws.title,
                "used_range": ws.dimensions,
                "rows": ws.max_row,
                "cols": ws.max_column,
                "merged": [str(r) for r in list(ws.merged_cells.ranges)[:200]],
                "tables": [{"name": t.name, "ref": t.ref} for t in ws.tables.values()],
                "has_pivot": bool(getattr(ws, "_pivots", None)),
                "pivots": _pivots(ws),
                "has_chart": bool(getattr(ws, "_charts", None)),
                "hidden": ws.sheet_state != "visible",
            }
        )
    for cs in wb.chartsheets:
        sheets.append(
            {
                "name": cs.title,
                "used_range": "",
                "rows": 0,
                "cols": 0,
                "merged": [],
                "tables": [],
                "has_pivot": False,
                "has_chart": True,
                "hidden": cs.sheet_state != "visible",
                "chartsheet": True,
            }
        )
    data = {"file_hash": sha256_file(real), "sheets": sheets, "has_macros": _zip_flags(real)}
    return Outcome(data=data)


def _formula_text(v: Any) -> str | None:
    """数式セルの数式（配列数式は ArrayFormula.text）。数式でなければ None。"""
    if isinstance(v, str):
        return v if v.startswith("=") else None
    text = getattr(v, "text", None)
    return text if isinstance(text, str) else None


def _resolve_sheet_range(wb: Workbook, ref: str) -> tuple[Worksheet, CellRange]:
    rng = parse_range(ref)
    name = (
        rng.sheet if rng.sheet is not None else wb.active.title if wb.active is not None else wb.sheetnames[0]
    )
    return sheet(wb, name), rng


@tool("xlsx.read_range")
def xlsx_read_range(ctx: ToolContext, args: dict[str, Any]) -> Outcome:
    path = args["path"]
    real = existing_file(path)
    mode = args.get("mode", "values")
    max_cells = int(args.get("max_cells", 2000))
    wb_v = open_book(path, real, data_only=True, read_only=True)
    try:
        ws_v, rng = _resolve_sheet_range(wb_v, args["range"])
        last_row, last_col = extent(ws_v)
        max_row = min(rng.max_row, max(last_row, rng.min_row))
        max_col = min(rng.max_col, max(last_col, rng.min_col))
        eff = CellRange(rng.sheet, rng.min_col, rng.min_row, max_col, max_row)
        header_row = args.get("header_row")
        header: list[str] | None = None
        first = eff.min_row
        if header_row is not None:
            hrow = next(
                ws_v.iter_rows(
                    min_row=header_row,
                    max_row=header_row,
                    min_col=eff.min_col,
                    max_col=eff.max_col,
                    values_only=True,
                ),
                (),
            )
            header = ["" if v is None else str(to_json(v)) for v in hrow]
            first = max(first, header_row + 1)
        values = (
            [
                [to_json(v) for v in row]
                for row in ws_v.iter_rows(
                    min_row=first,
                    max_row=eff.max_row,
                    min_col=eff.min_col,
                    max_col=eff.max_col,
                    values_only=True,
                )
            ]
            if first <= eff.max_row
            else []
        )
        formats: list[list[str]] | None = None
        if mode == "both":
            formats = (
                [
                    [str(getattr(c, "number_format", None) or "General") for c in row]
                    for row in ws_v.iter_rows(
                        min_row=first, max_row=eff.max_row, min_col=eff.min_col, max_col=eff.max_col
                    )
                ]
                if first <= eff.max_row
                else []
            )
        sheet_title = ws_v.title
    finally:
        wb_v.close()
    formulas: list[list[str | None]] | None = None
    if mode in ("formulas", "both"):
        wb_f = open_book(path, real, data_only=False, read_only=True)
        try:
            ws_f = sheet(wb_f, sheet_title)
            formulas = (
                [
                    [_formula_text(v) for v in row]
                    for row in ws_f.iter_rows(
                        min_row=first,
                        max_row=eff.max_row,
                        min_col=eff.min_col,
                        max_col=eff.max_col,
                        values_only=True,
                    )
                ]
                if first <= eff.max_row
                else []
            )
        finally:
            wb_f.close()
    items: list[Item] = []
    for i, row in enumerate(values):
        fields: dict[str, Any] = {"rows": row}
        if formulas is not None:
            fields["formulas"] = formulas[i] if i < len(formulas) else [None] * len(row)
        if formats is not None:
            fields["formats"] = formats[i] if i < len(formats) else ["General"] * len(row)
        items.append(Item(fields))
    base: dict[str, Any] = {"file_hash": sha256_file(real), "range": eff.qualified(sheet_title)}
    if header is not None:
        base["header"] = header
    keys = (
        ("rows",)
        + (("formulas",) if formulas is not None else ())
        + (("formats",) if formats is not None else ())
    )
    return paginate(ctx.store, "xlsx.read_range", base, keys, items, max_items=max(1, max_cells // eff.cols))


@tool("xlsx.find")
def xlsx_find(ctx: ToolContext, args: dict[str, Any]) -> Outcome:
    path = args["path"]
    real = existing_file(path)
    query = str(args["query"])
    exact = args.get("match", "contains") == "exact"
    limit = int(args.get("max_results", 100))
    q = query.casefold()
    hits: list[dict[str, Any]] = []
    wb = open_book(path, real, data_only=True, read_only=True)
    truncated = False
    try:
        sheets = [sheet(wb, args["sheet"])] if args.get("sheet") else wb.worksheets
        for ws in sheets:
            for row in ws.iter_rows():
                for c in row:
                    v = getattr(c, "value", None)
                    if v is None:
                        continue
                    j = to_json(v)
                    text = j["date"] if isinstance(j, dict) else str(j)
                    t = text.casefold()
                    if (t == q) if exact else (q in t):
                        if len(hits) >= limit:
                            truncated = True
                            break
                        value = j if not isinstance(j, str) else j[:200]
                        hits.append({"cell": f"{quote_sheet(ws.title)}!{c.coordinate}", "value": value})
                if truncated:
                    break
            if truncated:
                break
    finally:
        wb.close()
    out = Outcome(data={"hits": hits})
    if truncated:
        out.warnings.append(f"{limit} 件を超えたため打ち切りました。検索語を絞り込んでください")
    return out


# ---------------------------------------------------------------- 書き込み


@tool("xlsx.write_cells")
def xlsx_write_cells(ctx: ToolContext, args: dict[str, Any]) -> Outcome:
    path, real, wb = _load_for_write(args)
    ws = sheet(wb, args["sheet"])
    as_text: list[str] = []  # value に "=" で始まる文字列が来たセル（数式にせず文字列で書く）
    for cw in args["cells"]:
        ref_sheet, body = split_sheet(cw["ref"])
        if ref_sheet is not None and ref_sheet != ws.title:
            raise invalid(f"{cw['ref']} のシートが sheet 引数（{ws.title}）と違います")
        col, row = parse_cell(body)
        if "formula" in cw and "value" in cw:
            raise invalid(f"{cw['ref']}: value と formula は同時に指定できません")
        if "formula" in cw:
            cell = ws.cell(row=row, column=col)
            if isinstance(cell, MergedCell):
                raise invalid(f"{cell.coordinate} は結合セルの一部です。結合範囲の左上のセルに書いてください")
            cell.value = cw["formula"]
            if cw.get("number_format"):
                cell.number_format = cw["number_format"]
        elif "value" in cw:
            _write_value(ws, col, row, cw["value"], cw.get("number_format"))
            if isinstance(cw["value"], str) and cw["value"].startswith("="):
                as_text.append(cw["ref"])
        elif cw.get("number_format"):
            cell = ws.cell(row=row, column=col)
            cell.number_format = cw["number_format"]
    artifact, recalc, warnings = _save(path, real, wb, bool(args.get("recalc", True)))
    if as_text:
        # 写した値が勝手に数式として動かないよう、value は常に文字列で書く。数式にしたかった場合に気づけるよう知らせる
        warnings.append(
            f"{', '.join(as_text[:5])} の値は「=」で始まるため、数式ではなく文字列として書きました。"
            "数式にするなら value ではなく formula に指定してください"
        )
    return Outcome(
        data=_write_result(artifact, recalc, written=len(args["cells"])),
        artifacts=[artifact],
        warnings=warnings,
    )


@tool("xlsx.write_table")
def xlsx_write_table(ctx: ToolContext, args: dict[str, Any]) -> Outcome:
    path, real, wb = _load_for_write(args)
    ws = sheet(wb, args["sheet"])
    col0, row0 = parse_cell(args["top_left"])
    r = row0
    width = 0
    header = args.get("header")
    if header:
        for j, h in enumerate(header):
            _write_value(ws, col0 + j, r, h)
        width = len(header)
        r += 1
    for row in args["rows"]:
        for j, v in enumerate(row):
            _write_value(ws, col0 + j, r, v)
        width = max(width, len(row))
        r += 1
    last_row = max(row0, r - 1)
    rng = CellRange(None, col0, row0, col0 + max(width, 1) - 1, last_row)
    artifact, recalc, warnings = _save(path, real, wb, bool(args.get("recalc", True)))
    return Outcome(
        data=_write_result(artifact, recalc, range=rng.qualified(ws.title)),
        artifacts=[artifact],
        warnings=warnings,
    )


def _shift_tool(args: dict[str, Any], axis: str, delete: bool) -> Outcome:
    path, real, wb = _load_for_write(args)
    ws = sheet(wb, args["sheet"])
    at = int(args["at"]) if axis == "row" else col_index(str(args["at"]))
    warnings = apply_shift(wb, ws, Shift(ws.title, axis, at, int(args["count"]), delete))
    artifact, recalc, w2 = _save(path, real, wb, True)
    return Outcome(data=_write_result(artifact, recalc), artifacts=[artifact], warnings=warnings + w2)


@tool("xlsx.insert_rows")
def xlsx_insert_rows(ctx: ToolContext, args: dict[str, Any]) -> Outcome:
    return _shift_tool(args, "row", False)


@tool("xlsx.delete_rows")
def xlsx_delete_rows(ctx: ToolContext, args: dict[str, Any]) -> Outcome:
    return _shift_tool(args, "row", True)


@tool("xlsx.insert_cols")
def xlsx_insert_cols(ctx: ToolContext, args: dict[str, Any]) -> Outcome:
    return _shift_tool(args, "col", False)


@tool("xlsx.delete_cols")
def xlsx_delete_cols(ctx: ToolContext, args: dict[str, Any]) -> Outcome:
    return _shift_tool(args, "col", True)


@tool("xlsx.add_sheet")
def xlsx_add_sheet(ctx: ToolContext, args: dict[str, Any]) -> Outcome:
    path, real, wb = _load_for_write(args)
    name = args["name"]
    _check_sheet_name(name)
    if name in wb.sheetnames:
        raise invalid(f"シート {name} は既にあります")
    wb.create_sheet(name, args.get("position"))
    artifact, _, warnings = _save(path, real, wb, False)
    return Outcome(data={"artifact": artifact}, artifacts=[artifact], warnings=warnings)


@tool("xlsx.rename_sheet")
def xlsx_rename_sheet(ctx: ToolContext, args: dict[str, Any]) -> Outcome:
    path, real, wb = _load_for_write(args)
    old, new = args["name"], args["new_name"]
    _check_sheet_name(new)
    ws = sheet(wb, old)
    if new != old and new in wb.sheetnames:
        raise invalid(f"シート {new} は既にあります")
    ws.title = new
    rename_sheet_refs(wb, old, new)
    artifact, _, warnings = _save(path, real, wb, False)
    return Outcome(data={"artifact": artifact}, artifacts=[artifact], warnings=warnings)


@tool("xlsx.create")
def xlsx_create(ctx: ToolContext, args: dict[str, Any]) -> Outcome:
    path = args["path"]
    real = resolve(path)
    if ext(real) != ".xlsx":
        raise invalid("xlsx.create で作れるのは .xlsx だけです")
    ensure_absent(path, real, bool(args.get("overwrite", False)))
    wb = Workbook()
    default = wb.active
    names: set[str] = set()
    for spec in args["sheets"]:
        name = spec["name"]
        _check_sheet_name(name)
        if name in names:
            raise invalid(f"シート名 {name} が重複しています")
        names.add(name)
        ws = wb.create_sheet(name)
        cols = spec["columns"]
        for j, col in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=j, value=col["header"])
            cell.font = Font(bold=True)
            if col.get("width"):
                ws.column_dimensions[get_column_letter(j)].width = col["width"]
        rows = spec.get("rows", [])
        for i, row in enumerate(rows, start=2):
            for j, v in enumerate(row, start=1):
                fmt = cols[j - 1].get("number_format") if j <= len(cols) else None
                _write_value(ws, j, i, v, fmt)
        for f in spec.get("formulas", []):
            c, r = parse_cell(f["ref"])
            ws.cell(row=r, column=c, value=f["formula"])
            fmt = cols[c - 1].get("number_format") if c <= len(cols) else None
            if fmt:
                ws.cell(row=r, column=c).number_format = fmt
        if spec.get("freeze"):
            parse_cell(spec["freeze"])
            ws.freeze_panes = spec["freeze"]
        if spec.get("as_table"):
            headers = [c["header"] for c in cols]
            if len(set(headers)) != len(headers) or any(not h for h in headers):
                raise invalid(f"{name}: テーブルにするには見出しが空でなく重複しない必要があります")
            last = max(2, len(rows) + 1)
            ref = f"A1:{get_column_letter(len(cols))}{last}"
            table = Table(displayName=spec["as_table"]["name"], ref=ref)
            table.tableStyleInfo = TableStyleInfo(
                name=spec["as_table"].get("style", "TableStyleMedium2"), showRowStripes=True
            )
            ws.add_table(table)
    if default is not None:
        wb.remove(default)
    artifact, recalc, warnings = _save(path, real, wb, True)
    return Outcome(data=_write_result(artifact, recalc), artifacts=[artifact], warnings=warnings)


@tool("xlsx.recalc")
def xlsx_recalc(ctx: ToolContext, args: dict[str, Any]) -> Outcome:
    path = args["path"]
    real = existing_file(path)
    check_ext(path, real)
    formulas = rc.collect_formulas(real)
    if not formulas:
        return Outcome(data={"ok": True, "errors": []})
    if not rc.engine_available():
        return Outcome(
            data={"ok": False, "errors": [], "recalculated": False},
            warnings=[rc.NO_ENGINE_WARNING],
        )
    result: dict[str, Any] = {}

    def writer(tmp: Path) -> None:
        tmp.write_bytes(real.read_bytes())
        res = rc.recalc_path(tmp)
        result.update(res or {"ok": False, "errors": []})

    artifact = atomic_write(path, real, writer)
    return Outcome(data=result, artifacts=[artifact])


_CHARTS = {"bar": BarChart, "line": LineChart, "pie": PieChart, "scatter": ScatterChart}


@tool("xlsx.chart")
def xlsx_chart(ctx: ToolContext, args: dict[str, Any]) -> Outcome:
    path, real, wb = _load_for_write(args)
    ws = sheet(wb, args["sheet"])
    rng = parse_range(args["data_range"])
    src = sheet(wb, rng.sheet) if rng.sheet else ws
    kind = args["kind"]
    chart: Any = _CHARTS[kind]()
    if args.get("title"):
        chart.title = args["title"]
    if rng.rows < 2:
        raise invalid("data_range は見出し行とデータ行を含めて2行以上にしてください")
    if kind == "scatter":
        if rng.cols < 2:
            raise invalid("散布図には X 列と Y 列の2列以上が必要です")
        xs = Reference(src, min_col=rng.min_col, min_row=rng.min_row + 1, max_row=rng.max_row)
        for c in range(rng.min_col + 1, rng.max_col + 1):
            ys = Reference(src, min_col=c, min_row=rng.min_row, max_row=rng.max_row)
            chart.series.append(Series(ys, xs, title_from_data=True))
    else:
        has_cats = rng.cols >= 2
        data_min_col = rng.min_col + 1 if has_cats else rng.min_col
        data = Reference(
            src, min_col=data_min_col, min_row=rng.min_row, max_col=rng.max_col, max_row=rng.max_row
        )
        chart.add_data(data, titles_from_data=True)
        if has_cats:
            cats = Reference(src, min_col=rng.min_col, min_row=rng.min_row + 1, max_row=rng.max_row)
            chart.set_categories(cats)
    parse_cell(args["anchor"])
    ws.add_chart(chart, args["anchor"])
    artifact, _, warnings = _save(path, real, wb, False)
    return Outcome(data={"artifact": artifact}, artifacts=[artifact], warnings=warnings)
