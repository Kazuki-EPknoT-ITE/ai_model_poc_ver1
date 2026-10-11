"""data.transform（protocol §3.4、XLS-05）。AI は宣言的な spec を書き、計算はここで行う。"""

from __future__ import annotations

import datetime as dt
import re
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.utils.cell import get_column_letter
from openpyxl.worksheet.table import Table as XlTable
from openpyxl.worksheet.table import TableStyleInfo

from office_tools.envelope import Outcome
from office_tools.errors import invalid
from office_tools.fsutil import atomic_write, existing_file, resolve
from office_tools.registry import ToolContext, tool
from office_tools.scalars import to_json
from office_tools.tools.csvtool import read_csv_file, write_csv_bytes
from office_tools.transform.engine import Table, make_table, run_steps
from office_tools.transform.expr import to_decimal
from office_tools.xlsx.book import check_ext, extent, open_book, sheet
from office_tools.xlsx.refs import parse_range

PREVIEW_ROWS = 20
_ISO_DATE = re.compile(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}$")
_PLAIN_NUM = re.compile(r"^-?(0|[1-9]\d{0,2}(,\d{3})+|[1-9]\d*)(\.\d+)?$|^-?0\.\d+$")


def _from_cell(v: Any) -> Any:
    if isinstance(v, bool) or v is None:
        return v
    if isinstance(v, int | float):
        return to_decimal(v)
    if isinstance(v, dt.datetime):
        return v.date() if v.time() == dt.time(0, 0) else v
    return v


def _from_text(v: str) -> Any:
    """CSV の文字列を型付けする。先頭が 0 の番号（口座番号など）は文字列のまま。"""
    s = v.strip()
    if s == "":
        return None
    if _PLAIN_NUM.match(s):
        return Decimal(s.replace(",", ""))
    if _ISO_DATE.match(s):
        y, m, d = re.split(r"[-/]", s)
        try:
            return dt.date(int(y), int(m), int(d))
        except ValueError:
            return v
    return v


def _from_scalar(v: Any) -> Any:
    if isinstance(v, dict) and "date" in v:
        return dt.date.fromisoformat(str(v["date"]))
    if isinstance(v, bool) or v is None or isinstance(v, str):
        return v
    return to_decimal(v)


def _load_source(src: dict[str, Any]) -> Table:
    if "rows" in src:
        return make_table(src["header"], [[_from_scalar(v) for v in row] for row in src["rows"]])
    if "csv" in src:
        path = src["csv"]
        table = read_csv_file(path, existing_file(path))
        if not table.rows:
            return Table([])
        return make_table(table.rows[0], [[_from_text(v) for v in row] for row in table.rows[1:]])
    if "xlsx" in src:
        path = src["xlsx"]
        real = existing_file(path)
        wb = open_book(path, real, data_only=True, read_only=True)
        try:
            rng = parse_range(src["range"])
            ws = sheet(wb, rng.sheet) if rng.sheet else wb.worksheets[0]
            header_row = int(src.get("header_row", rng.min_row))
            max_row = min(rng.max_row, max(extent(ws)[0], header_row))
            header = next(
                ws.iter_rows(
                    min_row=header_row,
                    max_row=header_row,
                    min_col=rng.min_col,
                    max_col=rng.max_col,
                    values_only=True,
                ),
                (),
            )
            rows = [
                [_from_cell(v) for v in row]
                for row in ws.iter_rows(
                    min_row=header_row + 1,
                    max_row=max_row,
                    min_col=rng.min_col,
                    max_col=rng.max_col,
                    values_only=True,
                )
            ]
        finally:
            wb.close()
        return make_table(list(header), rows)
    raise invalid(
        "result_of は office-tools では解決できません。Runtime が rows と header に展開してから渡してください"
    )


def _write_xlsx(out: dict[str, Any], table: Table) -> dict[str, Any]:
    path = out["xlsx"]
    real = resolve(path)
    warnings: list[str] = []
    if real.exists():
        check_ext(path, real)
        wb = open_book(path, real)
    else:
        wb = Workbook()
        if wb.active is not None:
            wb.remove(wb.active)
    name = out["sheet"]
    top_left = out.get("top_left", "A1")
    if name in wb.sheetnames and top_left == "A1":
        idx = wb.sheetnames.index(name)
        wb.remove(wb[name])
        ws = wb.create_sheet(name, idx)
        warnings.append(f"既存のシート {name} を置き換えました")
    elif name in wb.sheetnames:
        ws = sheet(wb, name)
    else:
        ws = wb.create_sheet(name)
    rng = parse_range(top_left)
    c0, r0 = rng.min_col, rng.min_row
    for j, col in enumerate(table.columns):
        ws.cell(row=r0, column=c0 + j, value=col)
    for i, row in enumerate(table.matrix(), start=1):
        for j, v in enumerate(row):
            cell = ws.cell(row=r0 + i, column=c0 + j, value=v)
            if isinstance(v, dt.date):
                cell.number_format = "yyyy-mm-dd"
    if out.get("as_table") and table.columns:
        existing = {t for w in wb.worksheets for t in w.tables}
        n = 1
        while f"Result{n}" in existing:
            n += 1
        last_row = r0 + max(1, len(table.rows))
        ref = f"{get_column_letter(c0)}{r0}:{get_column_letter(c0 + len(table.columns) - 1)}{last_row}"
        t = XlTable(displayName=f"Result{n}", ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)

    def writer(tmp: Path) -> None:
        wb.save(tmp)

    artifact = atomic_write(path, real, writer)
    return {"artifact": artifact, "warnings": warnings}


@tool("data.transform")
def data_transform(ctx: ToolContext, args: dict[str, Any]) -> Outcome:
    tables: dict[str, Table] = {}
    for inp in args["inputs"]:
        if inp["name"] in tables:
            raise invalid(f"inputs の name {inp['name']} が重複しています")
        tables[inp["name"]] = _load_source(inp["source"])
    result = run_steps(tables, args["inputs"][0]["name"], args["steps"])
    data: dict[str, Any] = {
        "rows_out": len(result.rows),
        "columns": result.columns,
        "preview": [[to_json(v) for v in row] for row in result.matrix()[:PREVIEW_ROWS]],
    }
    outcome = Outcome(data=data)
    out = args["output"]
    if "csv" in out:
        path = out["csv"]
        payload = write_csv_bytes(result.columns, result.matrix(), "utf-8-sig")
        artifact = atomic_write(path, resolve(path), lambda tmp: tmp.write_bytes(payload))
        data["artifact"] = artifact
        outcome.artifacts.append(artifact)
    elif "xlsx" in out:
        written = _write_xlsx(out, result)
        data["artifact"] = written["artifact"]
        outcome.artifacts.append(written["artifact"])
        outcome.warnings.extend(written["warnings"])
    return outcome
