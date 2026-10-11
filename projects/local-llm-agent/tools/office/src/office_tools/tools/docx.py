"""docx.*（protocol §3.2、DOC-01〜06）。"""

from __future__ import annotations

import difflib
import hashlib
import io
import os
import re
import tempfile
import zipfile
from pathlib import Path
from typing import Any

import docx as python_docx
from docx.document import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from office_tools.docx.edit import BlockFactory, apply_ops, text_runs
from office_tools.docx.model import W_P, W_TC, Block, index_document, open_document
from office_tools.envelope import Item, Outcome, paginate, untrusted_block
from office_tools.errors import ToolFailure, invalid, unsupported
from office_tools.fsutil import (
    atomic_write,
    check_expect_hash,
    ensure_absent,
    existing_file,
    ext,
    normalize_arg,
    resolve,
    sha256_file,
)
from office_tools.registry import SamplingRequest, ToolContext, async_tool, tool

EXTRACT_DIR_ENV = "OFFICE_TOOLS_EXTRACT_DIR"


def _ref(path: str, ident: str) -> str:
    return f"{normalize_arg(path)}#{ident}"


def block_json(path: str, b: Block, include: set[str]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    """ブロック → data 側の JSON と untrusted。本文は untrusted に移し、data には text_ref だけを残す。"""
    out: dict[str, Any] = {"id": b.id, "kind": b.kind}
    unt: list[dict[str, Any]] = []
    if b.kind in ("heading", "paragraph"):
        out.update(b.info)
        out["chars"] = len(b.text)
    elif b.kind == "table":
        out.update(b.info)
    elif b.kind in ("header", "footer"):
        out["section"] = b.info["section"]
        if b.info.get("type", "default") != "default":
            out["type"] = b.info["type"]
    elif b.kind == "image":
        out["width_mm"] = b.info["width_mm"]
        out["height_mm"] = b.info["height_mm"]
        if b.info.get("alt"):
            ref = _ref(path, b.id) + "/alt"
            out["alt_ref"] = ref
            unt.append(untrusted_block(ref, str(b.info["alt"])))
    if b.text and b.kind != "image":
        ref = _ref(path, b.id)
        out["text_ref"] = ref
        unt.append(untrusted_block(ref, b.text))
    return out, unt


def _visible(b: Block, include: set[str]) -> bool:
    if b.kind in ("header", "footer"):
        return "headers" in include
    if b.kind == "image":
        return "images" in include
    return True


def _comments(doc: Document) -> list[dict[str, Any]]:
    try:
        items = list(doc.comments)
    except (KeyError, AttributeError, NotImplementedError):
        return []
    return [{"id": str(c.comment_id), "author": c.author, "text": c.text} for c in items]


@tool("docx.read")
def docx_read(ctx: ToolContext, args: dict[str, Any]) -> Outcome:
    path = args["path"]
    real = existing_file(path)
    doc = open_document(path, real)
    index = index_document(doc)
    include = set(args.get("include", []))
    blocks = [b for b in index.blocks if _visible(b, include)]
    if args.get("from_block"):
        start = next((i for i, b in enumerate(blocks) if b.id == args["from_block"]), None)
        if start is None:
            index.get(args["from_block"])  # BLOCK_NOT_FOUND
            raise invalid(f"{args['from_block']} は include の指定で読み取り対象から外れています")
        blocks = blocks[start:]
    by_block: dict[str, list[dict[str, Any]]] = {}
    if "comments" in include:
        for c in _comments(doc):
            bid = index.comment_anchor.get(c["id"])
            if bid:
                by_block.setdefault(bid, []).append(c)
    items: list[Item] = []
    for b in blocks:
        out, unt = block_json(path, b, include)
        for c in by_block.get(b.id, []):
            ref = _ref(path, f"comment:{c['id']}")
            out.setdefault("comments", []).append({"id": c["id"], "author": c["author"], "text_ref": ref})
            unt.append(untrusted_block(ref, c["text"]))
        items.append(Item({"blocks": out}, unt))
    base = {
        "file_hash": sha256_file(real),
        "total_blocks": len([b for b in index.blocks if _visible(b, include)]),
    }
    return paginate(
        ctx.store, "docx.read", base, ("blocks",), items, max_items=int(args.get("max_blocks", 200))
    )


def _pages_estimate(real: Path, chars: int, breaks: int) -> int:
    try:
        with zipfile.ZipFile(real) as z:
            if "docProps/app.xml" in z.namelist():
                m = re.search(rb"<(?:\w+:)?Pages>(\d+)</", z.read("docProps/app.xml"))
                if m:
                    return max(1, int(m.group(1)))
    except (OSError, zipfile.BadZipFile):
        pass
    return max(1, breaks + 1, chars // 1000 + 1)


@tool("docx.outline")
def docx_outline(ctx: ToolContext, args: dict[str, Any]) -> Outcome:
    path = args["path"]
    real = existing_file(path)
    doc = open_document(path, real)
    index = index_document(doc)
    headings = []
    unt = []
    for b in index.blocks:
        if b.kind == "heading":
            ref = _ref(path, b.id)
            headings.append({"id": b.id, "level": b.info["level"], "text_ref": ref})
            unt.append(untrusted_block(ref, b.text))
    chars = sum(len(b.text) for b in index.blocks if b.kind in ("heading", "paragraph", "table"))
    breaks = sum(1 for b in index.blocks if b.kind == "page_break")
    data = {
        "file_hash": sha256_file(real),
        "headings": headings,
        "tables": sum(1 for b in index.blocks if b.kind == "table"),
        "pages_estimate": _pages_estimate(real, chars, breaks),
    }
    return Outcome(data=data, untrusted=unt)


def _save_doc(path: str, real: Path, doc: Document) -> dict[str, Any]:
    return atomic_write(path, real, lambda tmp: doc.save(str(tmp)))


@tool("docx.create")
def docx_create(ctx: ToolContext, args: dict[str, Any]) -> Outcome:
    path = args["path"]
    real = resolve(path)
    if ext(real) != ".docx":
        raise invalid("docx.create で作れるのは .docx だけです")
    ensure_absent(path, real, bool(args.get("overwrite", False)))
    if args.get("template"):
        tpath = args["template"]
        treal = existing_file(tpath)
        doc = open_document(tpath, treal)
        body = doc.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
    else:
        doc = python_docx.Document()
    factory = BlockFactory(doc)
    body = doc.element.body
    sect = body.find(qn("w:sectPr"))
    for nb in args["blocks"]:
        el = factory.build(nb)
        if sect is not None:
            sect.addprevious(el)
        else:
            body.append(el)
    artifact = _save_doc(path, real, doc)
    return Outcome(data={"artifact": artifact}, artifacts=[artifact], warnings=factory.warnings)


@tool("docx.edit")
def docx_edit(ctx: ToolContext, args: dict[str, Any]) -> Outcome:
    path = args["path"]
    real = existing_file(path)
    check_expect_hash(real, args.get("expect_hash"))
    doc = open_document(path, real, for_edit=True)
    index = index_document(doc)
    results, applied, warnings = apply_ops(
        doc, index, args["ops"], bool(args.get("track_changes", False)), args.get("author", "AI")
    )
    if applied == 0:
        raise ToolFailure(
            "INVALID_ARGS",
            "どの op も適用できませんでした。results を確認し、docx.read で読み直してください",
            detail={"results": results},
        )
    artifact = _save_doc(path, real, doc)
    failed_ops = [r["op_index"] for r in results if not r["ok"]]
    if failed_ops:
        warnings.append(f"op {failed_ops} は適用できませんでした（他の op は保存済み）")
    return Outcome(
        data={"artifact": artifact, "applied": applied, "results": results},
        artifacts=[artifact],
        warnings=warnings,
    )


# ---------------------------------------------------------------- diff


def _units(path: str) -> list[tuple[str, str]]:
    real = existing_file(path)
    doc = open_document(path, real)
    return [(b.id, b.text) for b in index_document(doc).blocks if b.kind in ("heading", "paragraph", "table")]


def _inline_diff(a: str, b: str) -> str:
    sm = difflib.SequenceMatcher(None, a, b, autojunk=False)
    out = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            seg = a[i1:i2]
            out.append(seg if len(seg) <= 40 else seg[:15] + "…" + seg[-15:])
        else:
            if i2 > i1:
                out.append("[-" + a[i1:i2] + "-]")
            if j2 > j1:
                out.append("{+" + b[j1:j2] + "+}")
    return "".join(out)


@tool("docx.diff")
def docx_diff(ctx: ToolContext, args: dict[str, Any]) -> Outcome:
    a_path, b_path = args["a"], args["b"]
    ua, ub = _units(a_path), _units(b_path)
    char = args.get("granularity", "char") == "char"
    sm = difflib.SequenceMatcher(None, [t for _, t in ua], [t for _, t in ub], autojunk=False)
    items: list[Item] = []

    def add(kind: str, ai: int | None, bi: int | None) -> None:
        entry: dict[str, Any] = {"kind": kind}
        if ai is not None:
            entry["a_id"] = ua[ai][0]
        if bi is not None:
            entry["b_id"] = ub[bi][0]
        before = ua[ai][1] if ai is not None else ""
        after = ub[bi][1] if bi is not None else ""
        if kind == "modify":
            text = _inline_diff(before, after) if char else f"- {before}\n+ {after}"
        elif kind == "insert":
            text = f"+ {after}"
        else:
            text = f"- {before}"
        ref = _ref(b_path, ub[bi][0]) if bi is not None else _ref(a_path, ua[ai][0] if ai is not None else "")
        entry["text_ref"] = ref
        items.append(Item({"changes": entry}, [untrusted_block(ref, text)]))

    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            continue
        if tag == "delete":
            for i in range(i1, i2):
                add("delete", i, None)
        elif tag == "insert":
            for j in range(j1, j2):
                add("insert", None, j)
        else:
            n = min(i2 - i1, j2 - j1)
            for k in range(n):
                add("modify", i1 + k, j1 + k)
            for i in range(i1 + n, i2):
                add("delete", i, None)
            for j in range(j1 + n, j2):
                add("insert", None, j)
    base: dict[str, Any] = {"total_changes": len(items)}
    return paginate(ctx.store, "docx.diff", base, ("changes",), items, max_items=500)


# ---------------------------------------------------------------- comments


def _anchor_paragraph(doc: Document, path: str, block_id: str) -> Paragraph:
    index = index_document(doc)
    if "/" in block_id:
        _, tc, _ = index.cell(block_id)
        p_el = tc.find(W_P)
    else:
        b = index.get(block_id)
        if b.kind in ("paragraph", "heading", "page_break"):
            p_el = b.para
        elif b.kind == "image":
            p_el = b.para
        elif b.kind == "table":
            first_tc = next(b.el.iter(W_TC), None)
            p_el = first_tc.find(W_P) if first_tc is not None else None
        else:
            raise invalid(f"{block_id} にはコメントを付けられません")
    if p_el is None:
        raise invalid(f"{block_id} にはコメントを付けられません")
    return Paragraph(p_el, doc._body)


@tool("docx.comments")
def docx_comments(ctx: ToolContext, args: dict[str, Any]) -> Outcome:
    path = args["path"]
    real = existing_file(path)
    adds = args.get("add") or []
    doc = open_document(path, real, for_edit=bool(adds))
    artifacts = []
    for item in adds:
        para = _anchor_paragraph(doc, path, item["block_id"])
        runs = [r for r in para.runs if r._r in text_runs(para._p)] or [para.add_run("")]
        doc.add_comment(runs, text=item["text"], author="AI", initials="AI")
    data: dict[str, Any] = {}
    if adds:
        artifact = _save_doc(path, real, doc)
        artifacts.append(artifact)
        data["artifact"] = artifact
        doc = open_document(path, real)
    index = index_document(doc)
    comments = []
    unt = []
    for c in _comments(doc):
        ref = _ref(path, f"comment:{c['id']}")
        comments.append(
            {
                "id": c["id"],
                "block_id": index.comment_anchor.get(c["id"]),
                "author": c["author"],
                "text_ref": ref,
            }
        )
        unt.append(untrusted_block(ref, c["text"]))
    data["comments"] = comments
    return Outcome(data=data, untrusted=unt, artifacts=artifacts)


# ---------------------------------------------------------------- images


def _extract_dir(real: Path) -> Path:
    base = Path(os.environ.get(EXTRACT_DIR_ENV) or Path(tempfile.gettempdir()) / "office-tools")
    digest = hashlib.sha256(str(real).encode("utf-8")).hexdigest()[:12]
    out = base / f"docx-{digest}"
    out.mkdir(parents=True, exist_ok=True)
    return out


def _to_png(blob: bytes) -> bytes | None:
    try:
        from PIL import Image

        with Image.open(io.BytesIO(blob)) as img:
            buf = io.BytesIO()
            img.convert("RGB").save(buf, format="PNG")
            return buf.getvalue()
    except Exception:
        return None


def _images(path: str, real: Path) -> list[tuple[str, str, bytes]]:
    e = ext(real)
    if e not in (".docx", ".docm", ".dotx", ".dotm"):
        raise unsupported(path, f"{e} は Word 文書として扱えません")
    doc = open_document(path, real)
    index = index_document(doc)
    out = []
    for b in index.blocks:
        if b.kind != "image":
            continue
        blip = next(b.el.iter(qn("a:blip")), None)
        rid = blip.get(qn("r:embed")) if blip is not None else None
        part = doc.part.related_parts.get(rid) if rid else None
        if part is None:
            continue
        suffix = Path(str(part.partname)).suffix or ".bin"
        out.append((b.id, suffix, bytes(part.blob)))
    return out


@async_tool("docx.images")
async def docx_images(ctx: ToolContext, args: dict[str, Any]) -> Outcome:
    import anyio.to_thread

    path = args["path"]
    real = existing_file(path)
    found = await anyio.to_thread.run_sync(lambda: _images(path, real))
    out_dir = _extract_dir(real)
    images = []
    unt = []
    for bid, suffix, blob in found:
        target = out_dir / f"{bid.replace(':', '_')}{suffix}"
        target.write_bytes(blob)
        entry: dict[str, Any] = {"id": bid, "path_extracted": str(target)}
        if args.get("describe"):
            png = _to_png(blob)
            if png is not None:
                text = await ctx.sample(
                    SamplingRequest(
                        purpose="image_describe", image_png=png, metadata={"block_id": bid}, max_tokens=500
                    )
                )
                ref = _ref(path, bid) + "/description"
                entry["description_ref"] = ref
                unt.append(untrusted_block(ref, text, source="ocr"))
        images.append(entry)
    return Outcome(data={"images": images}, untrusted=unt)
