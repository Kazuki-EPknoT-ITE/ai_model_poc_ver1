"""ブックを開く・保存する（.xlsm は VBA を保持、.xls は非対応）。"""

from __future__ import annotations

import warnings
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chartsheet import Chartsheet
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from office_tools.errors import corrupt, invalid, locked, not_found, unsupported
from office_tools.fsutil import ext

XLSX_EXTS = {".xlsx", ".xlsm", ".xltx", ".xltm"}


def check_ext(path: str, real: Path) -> None:
    e = ext(real)
    if e == ".xls":
        raise unsupported(path, ".xls は convert.to_xlsx で .xlsx に変換してから扱ってください")
    if e not in XLSX_EXTS:
        raise unsupported(
            path, f"{e or '拡張子なし'} は Excel ブックとして扱えません（.xlsx / .xlsm に対応）"
        )


def is_macro(real: Path) -> bool:
    return ext(real) in (".xlsm", ".xltm")


def open_book(path: str, real: Path, *, data_only: bool = False, read_only: bool = False) -> Workbook:
    check_ext(path, real)
    if not real.exists():
        raise not_found(path)
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            return load_workbook(
                real,
                data_only=data_only,
                read_only=read_only,
                keep_vba=is_macro(real) and not read_only,
                keep_links=True,
                rich_text=not read_only,
            )
    except PermissionError as e:
        raise locked(path) from e
    except (zipfile.BadZipFile, InvalidFileException, KeyError, ValueError, OSError, TypeError) as e:
        raise corrupt(path, type(e).__name__) from e


def sheet(wb: Workbook, name: str) -> Worksheet:
    if name not in wb.sheetnames:
        raise invalid(
            f"シート {name} はありません。xlsx.info でシート名を確認してください",
            sheets=wb.sheetnames[:50],
        )
    ws = wb[name]
    if isinstance(ws, Chartsheet):
        raise invalid(f"{name} はワークシートではありません（グラフシートなど）")
    return ws


def has_formulas(wb: Workbook) -> bool:
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == "f":
                    return True
    return False


def mark_full_calc(wb: Workbook) -> None:
    """キャッシュ値の無い数式があるので、Excel で開いたときに全再計算させる。"""
    calc: Any = wb.calculation
    if calc is not None:
        calc.fullCalcOnLoad = True


def extent(ws: Any) -> tuple[int, int]:
    """使用範囲の最終行・最終列。寸法の記録が無いブック（他のツールで作ったものなど）は走査して求める。"""
    if ws.max_row is None or ws.max_column is None:
        calc = getattr(ws, "calculate_dimension", None)
        if calc is not None:
            calc(force=True)
    return int(ws.max_row or 1), int(ws.max_column or 1)
