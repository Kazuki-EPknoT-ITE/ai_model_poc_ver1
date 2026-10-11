"""セル参照（`Sheet1!B12`, `'送金 明細'!A1:F200`）の解析。"""

from __future__ import annotations

import re
from dataclasses import dataclass

from openpyxl.utils.cell import column_index_from_string, get_column_letter

from office_tools.errors import invalid

_CELL = re.compile(r"^\$?([A-Za-z]{1,3})\$?([0-9]{1,7})$")


@dataclass(frozen=True)
class CellRange:
    sheet: str | None
    min_col: int
    min_row: int
    max_col: int
    max_row: int

    @property
    def ref(self) -> str:
        a = f"{get_column_letter(self.min_col)}{self.min_row}"
        b = f"{get_column_letter(self.max_col)}{self.max_row}"
        return a if a == b else f"{a}:{b}"

    def qualified(self, sheet: str) -> str:
        return f"{quote_sheet(sheet)}!{self.ref}"

    @property
    def rows(self) -> int:
        return self.max_row - self.min_row + 1

    @property
    def cols(self) -> int:
        return self.max_col - self.min_col + 1


def quote_sheet(name: str) -> str:
    if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_.]*", name) and not _CELL.match(name):
        return name
    return "'" + name.replace("'", "''") + "'"


def split_sheet(ref: str) -> tuple[str | None, str]:
    ref = ref.strip()
    if ref.startswith("'"):
        i = 1
        name = []
        while i < len(ref):
            ch = ref[i]
            if ch == "'":
                if i + 1 < len(ref) and ref[i + 1] == "'":
                    name.append("'")
                    i += 2
                    continue
                break
            name.append(ch)
            i += 1
        rest = ref[i + 1 :]
        if not rest.startswith("!"):
            raise invalid(f"参照 {ref} の形式が不正です。'シート名'!A1 の形式で指定してください")
        return "".join(name), rest[1:]
    if "!" in ref:
        sheet, _, cell = ref.rpartition("!")
        return sheet, cell
    return None, ref


def parse_cell(text: str) -> tuple[int, int]:
    """'B12' → (col, row)。"""
    m = _CELL.match(text.strip())
    if not m:
        raise invalid(f"セル参照 {text} の形式が不正です。B12 の形式で指定してください")
    col = column_index_from_string(m.group(1).upper())
    row = int(m.group(2))
    if col > 16384 or row < 1 or row > 1048576:
        raise invalid(f"セル参照 {text} がシートの範囲外です")
    return col, row


def parse_range(ref: str) -> CellRange:
    sheet, body = split_sheet(ref)
    if ":" in body:
        a, _, b = body.partition(":")
        c1, r1 = parse_cell(a)
        c2, r2 = parse_cell(b)
    else:
        c1, r1 = parse_cell(body)
        c2, r2 = c1, r1
    return CellRange(sheet, min(c1, c2), min(r1, r2), max(c1, c2), max(r1, r2))


def col_index(letters: str) -> int:
    return column_index_from_string(letters.upper())
