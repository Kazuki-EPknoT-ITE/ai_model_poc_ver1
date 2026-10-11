"""行・列の挿入/削除とシート名変更に伴う参照の付け替え。

openpyxl の insert_rows などはセルを動かすだけで、数式・結合セル・入力規則・条件付き書式・
テーブル・名前定義の参照を更新しない。Excel と同じ意味で付け替える。
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formula.tokenizer import Token, Tokenizer, TokenizerError
from openpyxl.utils.cell import column_index_from_string, get_column_letter
from openpyxl.worksheet.cell_range import CellRange, MultiCellRange
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.worksheet import Worksheet

from office_tools.xlsx.refs import quote_sheet

_CELL = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?)([0-9]+)$")
_COL = re.compile(r"^(\$?)([A-Za-z]{1,3})$")
_ROW = re.compile(r"^(\$?)([0-9]+)$")


@dataclass(frozen=True)
class Shift:
    sheet: str
    axis: str  # "row" | "col"
    at: int
    count: int
    delete: bool

    def index(self, i: int) -> int | None:
        if not self.delete:
            return i + self.count if i >= self.at else i
        if i < self.at:
            return i
        if i >= self.at + self.count:
            return i - self.count
        return None

    def span(self, lo: int, hi: int) -> tuple[int, int] | None:
        if not self.delete:
            return (lo + self.count if lo >= self.at else lo, hi + self.count if hi >= self.at else hi)
        last = self.at + self.count - 1
        new_lo = lo if lo < self.at else (self.at if lo <= last else lo - self.count)
        new_hi = hi if hi < self.at else (self.at - 1 if hi <= last else hi - self.count)
        if new_hi < new_lo:
            return None
        return new_lo, new_hi


# ---------------------------------------------------------------- 参照文字列


def split_ref(text: str) -> tuple[str | None, str, str]:
    """'Sheet 1'!A1:B2 → (シート名, 元のシート接頭辞, 本体)。"""
    if text.startswith("'"):
        i = 1
        buf = []
        while i < len(text):
            if text[i] == "'":
                if i + 1 < len(text) and text[i + 1] == "'":
                    buf.append("'")
                    i += 2
                    continue
                break
            buf.append(text[i])
            i += 1
        if text[i + 1 : i + 2] == "!":
            return "".join(buf), text[: i + 2], text[i + 2 :]
        return None, "", text
    if "!" in text:
        prefix, _, body = text.rpartition("!")
        return prefix, prefix + "!", body
    return None, "", text


def shift_body(body: str, s: Shift) -> str | None:
    """参照本体を付け替える。参照でない（名前など）ならそのまま返す。"""
    parts = body.split(":")
    if len(parts) > 2:
        return body
    if s.axis == "row":
        cells = [_CELL.match(p) for p in parts]
        if all(cells):
            ms = [m for m in cells if m]
            if len(ms) == 1:
                m = ms[0]
                r = s.index(int(m.group(4)))
                return None if r is None else f"{m.group(1)}{m.group(2)}{m.group(3)}{r}"
            span = s.span(int(ms[0].group(4)), int(ms[1].group(4)))
            if span is None:
                return None
            a, b = ms
            return (
                f"{a.group(1)}{a.group(2)}{a.group(3)}{span[0]}:{b.group(1)}{b.group(2)}{b.group(3)}{span[1]}"
            )
        rows = [_ROW.match(p) for p in parts]
        if len(parts) == 2 and all(rows):
            ra, rb = (m for m in rows if m)
            span = s.span(int(ra.group(2)), int(rb.group(2)))
            if span is None:
                return None
            return f"{ra.group(1)}{span[0]}:{rb.group(1)}{span[1]}"
        return body
    # 列
    cells = [_CELL.match(p) for p in parts]
    if all(cells):
        ms = [m for m in cells if m]
        cols = [column_index_from_string(m.group(2).upper()) for m in ms]
        if len(ms) == 1:
            c = s.index(cols[0])
            m = ms[0]
            return None if c is None else f"{m.group(1)}{get_column_letter(c)}{m.group(3)}{m.group(4)}"
        span = s.span(cols[0], cols[1])
        if span is None:
            return None
        a, b = ms
        return (
            f"{a.group(1)}{get_column_letter(span[0])}{a.group(3)}{a.group(4)}:"
            f"{b.group(1)}{get_column_letter(span[1])}{b.group(3)}{b.group(4)}"
        )
    colm = [_COL.match(p) for p in parts]
    if len(parts) == 2 and all(colm):
        ca, cb = (m for m in colm if m)
        span = s.span(
            column_index_from_string(ca.group(2).upper()), column_index_from_string(cb.group(2).upper())
        )
        if span is None:
            return None
        return f"{ca.group(1)}{get_column_letter(span[0])}:{cb.group(1)}{get_column_letter(span[1])}"
    return body


class Rewriter:
    """参照の付け替え規則。"""

    def sheet(self, name: str) -> str:
        return name

    def body(self, sheet: str | None, body: str) -> str | None:
        return body


class ShiftRewriter(Rewriter):
    def __init__(self, s: Shift) -> None:
        self.s = s

    def body(self, sheet: str | None, body: str) -> str | None:
        if sheet != self.s.sheet:
            return body
        return shift_body(body, self.s)


class RenameRewriter(Rewriter):
    def __init__(self, old: str, new: str) -> None:
        self.old = old
        self.new = new

    def sheet(self, name: str) -> str:
        return self.new if name == self.old else name


def rewrite_formula(formula: str, host_sheet: str | None, rw: Rewriter) -> str:
    """数式中の範囲参照を付け替える（"=..." の形で渡す）。"""
    if not formula.startswith("="):
        return formula
    try:
        tok = Tokenizer(formula)
    except TokenizerError:
        return formula
    changed = False
    for t in tok.items:
        if t.type != Token.OPERAND or t.subtype != Token.RANGE:
            continue
        sheet, prefix, body = split_ref(t.value)
        if prefix.startswith("["):
            continue
        if prefix and sheet is not None:
            new = rw.body(sheet, body)
            new_sheet = rw.sheet(sheet)
            new_prefix = prefix if new_sheet == sheet else quote_sheet(new_sheet) + "!"
            value = "#REF!" if new is None else new_prefix + new
        else:
            new = rw.body(host_sheet, body) if host_sheet is not None else body
            value = "#REF!" if new is None else new
        if value != t.value:
            t.value = value
            changed = True
    return "=" + "".join(t.value for t in tok.items) if changed else formula


def shift_sqref(sqref: str, s: Shift) -> str:
    out = []
    for part in sqref.split():
        new = shift_body(part, s)
        if new is not None:
            out.append(new)
    return " ".join(out)


# ---------------------------------------------------------------- ブック全体


def _rewrite_all_formulas(wb: Workbook, fn: Rewriter) -> None:
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.data_type != "f":
                    continue
                v: Any = c.value
                if isinstance(v, str):
                    c.value = rewrite_formula(v, ws.title, fn)
                elif isinstance(v, ArrayFormula) and v.text:
                    v.text = rewrite_formula(v.text, ws.title, fn)
        for dv in ws.data_validations.dataValidation:
            for attr in ("formula1", "formula2"):
                f = getattr(dv, attr)
                if f:
                    new = rewrite_formula("=" + f, ws.title, fn)[1:]
                    setattr(dv, attr, new)
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                if rule.formula:
                    rule.formula = [rewrite_formula("=" + f, ws.title, fn)[1:] for f in rule.formula]
    for dn in list(wb.defined_names.values()):
        if dn.attr_text:
            dn.attr_text = rewrite_formula("=" + dn.attr_text, None, fn)[1:]
    for ws in wb.worksheets:
        for dn in list(ws.defined_names.values()):
            if dn.attr_text:
                dn.attr_text = rewrite_formula("=" + dn.attr_text, None, fn)[1:]


def rename_sheet_refs(wb: Workbook, old: str, new: str) -> None:
    _rewrite_all_formulas(wb, RenameRewriter(old, new))


def apply_shift(wb: Workbook, ws: Worksheet, s: Shift) -> list[str]:
    """セルを動かし、参照を付け替える。警告の一覧を返す。"""
    warnings: list[str] = []
    if s.axis == "row":
        (ws.delete_rows if s.delete else ws.insert_rows)(s.at, s.count)
    else:
        (ws.delete_cols if s.delete else ws.insert_cols)(s.at, s.count)
    _rewrite_all_formulas(wb, ShiftRewriter(s))
    # 配列数式の範囲
    for row in ws.iter_rows():
        for c in row:
            v: Any = c.value
            if isinstance(v, ArrayFormula) and v.ref:
                new = shift_body(v.ref, s)
                if new:
                    v.ref = new
            if c.hyperlink is not None:
                c.hyperlink.ref = c.coordinate
    # 結合セル
    for mcr in list(ws.merged_cells.ranges):
        new = shift_body(mcr.coord, s)
        if new is None:
            ws.merged_cells.remove(mcr)
            continue
        lo, _, hi = new.partition(":")
        cr = CellRange(new if hi else f"{lo}:{lo}")
        mcr.min_row, mcr.max_row, mcr.min_col, mcr.max_col = cr.min_row, cr.max_row, cr.min_col, cr.max_col
    # 入力規則の範囲
    for dv in list(ws.data_validations.dataValidation):
        new_sq = shift_sqref(str(dv.sqref), s)
        if new_sq:
            dv.sqref = MultiCellRange(new_sq)
        else:
            ws.data_validations.dataValidation.remove(dv)
    # 条件付き書式の範囲
    old_cf = ws.conditional_formatting
    new_cf = ConditionalFormattingList()
    for cf in old_cf:
        new_sq = shift_sqref(str(cf.sqref), s)
        if not new_sq:
            continue
        for rule in cf.rules:
            new_cf.add(new_sq, rule)
    ws.conditional_formatting = new_cf
    # テーブル
    for name in list(ws.tables.keys()):
        table = ws.tables[name]
        new = shift_body(table.ref, s)
        if new is None or ":" not in new:
            del ws.tables[name]
            warnings.append(f"テーブル {name} は範囲が削除されたため取り除きました")
            continue
        table.ref = new
        if table.autoFilter is not None:
            table.autoFilter.ref = new
    if ws.auto_filter.ref:
        new = shift_body(ws.auto_filter.ref, s)
        ws.auto_filter.ref = new
    if ws.freeze_panes:
        new = shift_body(ws.freeze_panes, s)
        if new:
            ws.freeze_panes = new
    # 行の高さ・列の幅
    if s.axis == "row":
        dims = dict(ws.row_dimensions.items())
        ws.row_dimensions.clear()
        for idx, dim in sorted(dims.items()):
            new_idx = s.index(idx)
            if new_idx is not None:
                dim.index = new_idx
                ws.row_dimensions[new_idx] = dim
    else:
        cdims = dict(ws.column_dimensions.items())
        ws.column_dimensions.clear()
        for letter, cdim in cdims.items():
            idx = column_index_from_string(letter)
            new_idx = s.index(idx)
            if new_idx is None:
                continue
            new_letter = get_column_letter(new_idx)
            cdim.index = new_letter
            span = s.span(cdim.min or idx, cdim.max or idx)
            if span:
                cdim.min, cdim.max = span
            ws.column_dimensions[new_letter] = cdim
    if any(getattr(w, "_charts", None) for w in wb.worksheets):
        warnings.append("グラフの参照範囲は自動では更新していません。必要なら確認してください")
    return warnings
