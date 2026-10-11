"""再計算（XLS-04）。

LibreOffice で一時コピーを再計算し、計算結果を元のブックの数式セルのキャッシュ値（`<v>`）
として書き込む。ブックそのものは LibreOffice で保存し直さない（書式・VBA・入力規則を
LibreOffice の往復で崩さないため）。エラー値のセルを一覧にして返す。
"""

from __future__ import annotations

import datetime as dt
import os
import posixpath
import shutil
import tempfile
import warnings
import zipfile
from pathlib import Path
from typing import Any

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel

from office_tools import libreoffice
from office_tools.xlsx.refs import quote_sheet

NO_ENGINE_WARNING = "再計算エンジン(LibreOffice)が見つからないため再計算していません"
ERROR_VALUES = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"

Formulas = dict[str, dict[str, str]]
"""シート名 → セル座標 → 数式。"""
Values = dict[str, dict[str, Any]]


def engine_available() -> bool:
    return libreoffice.find_soffice() is not None


def collect_formulas(path: Path) -> Formulas:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(path, read_only=True, data_only=False, keep_links=True)
    out: Formulas = {}
    try:
        for ws in wb.worksheets:
            cells: dict[str, str] = {}
            for row in ws.iter_rows():
                for c in row:
                    v = getattr(c, "value", None)
                    if getattr(c, "data_type", None) == "f":
                        text = v if isinstance(v, str) else getattr(v, "text", None) or str(v)
                        cells[c.coordinate] = text
            if cells:
                out[ws.title] = cells
    finally:
        wb.close()
    return out


def compute_values(path: Path, wanted: Formulas) -> Values:
    """LibreOffice で再計算した値を返す（wanted のセルだけ）。"""
    with tempfile.TemporaryDirectory(prefix="lla-recalc-") as td:
        src = Path(td) / f"book{path.suffix.lower()}"
        shutil.copyfile(path, src)
        out_dir = Path(td) / "out"
        out_dir.mkdir()
        out = libreoffice.convert(src, "xlsx", out_dir)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = load_workbook(out, data_only=True)
        result: Values = {}
        for name, cells in wanted.items():
            if name not in wb.sheetnames:
                continue
            ws = wb[name]
            result[name] = {coord: ws[coord].value for coord in cells}
        wb.close()
        return result


def find_errors(formulas: Formulas, values: Values) -> list[dict[str, Any]]:
    errors: list[dict[str, Any]] = []
    for sheet_name, cells in formulas.items():
        for coord, formula in cells.items():
            v = values.get(sheet_name, {}).get(coord)
            if isinstance(v, str) and v in ERROR_VALUES:
                errors.append({"cell": f"{quote_sheet(sheet_name)}!{coord}", "value": v, "formula": formula})
    return errors


def _sheet_parts(z: zipfile.ZipFile) -> dict[str, str]:
    """シート名 → zip 内のパス。"""
    wb_xml = etree.fromstring(z.read("xl/workbook.xml"))
    rels_xml = etree.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    targets = {}
    for rel in rels_xml.findall(f"{{{NS_PKG_REL}}}Relationship"):
        target = rel.get("Target", "")
        path = (
            target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
        )
        targets[rel.get("Id")] = path
    out = {}
    for s in wb_xml.iter(f"{{{NS_MAIN}}}sheet"):
        rid = s.get(f"{{{NS_REL}}}id")
        if rid in targets:
            out[s.get("name", "")] = targets[rid]
    return out


def _drop_type(c: etree._Element) -> None:
    if "t" in c.attrib:
        del c.attrib["t"]


def _set_cached(c: etree._Element, value: Any) -> None:
    v = c.find(f"{{{NS_MAIN}}}v")
    if v is None:
        v = etree.SubElement(c, f"{{{NS_MAIN}}}v")
    if isinstance(value, bool):
        c.set("t", "b")
        v.text = "1" if value else "0"
    elif isinstance(value, str):
        c.set("t", "e" if value in ERROR_VALUES else "str")
        v.text = value
    elif isinstance(value, dt.datetime | dt.date | dt.time):
        _drop_type(c)
        v.text = repr(float(to_excel(value)))  # type: ignore[no-untyped-call]
    elif isinstance(value, int | float):
        _drop_type(c)
        v.text = repr(value) if isinstance(value, float) else str(value)
    else:
        c.remove(v)


def inject_cached_values(path: Path, values: Values) -> None:
    """数式セルにキャッシュ値を書き込む（それ以外の XML は変えない）。"""
    with zipfile.ZipFile(path) as z:
        parts = _sheet_parts(z)
        replaced: dict[str, bytes] = {}
        for sheet_name, cells in values.items():
            part = parts.get(sheet_name)
            if part is None or part not in z.namelist():
                continue
            root = etree.fromstring(z.read(part))
            changed = False
            for c in root.iter(f"{{{NS_MAIN}}}c"):
                coord = c.get("r")
                if coord in cells and c.find(f"{{{NS_MAIN}}}f") is not None and cells[coord] is not None:
                    _set_cached(c, cells[coord])
                    changed = True
            if changed:
                replaced[part] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        if not replaced:
            return
        tmp = path.with_name(path.name + ".inj")
        with zipfile.ZipFile(tmp, "w") as out:
            for info in z.infolist():
                data = replaced.get(info.filename)
                out.writestr(info, data if data is not None else z.read(info.filename))
    os.replace(tmp, path)


def recalc_path(path: Path) -> dict[str, Any] | None:
    """path を再計算してキャッシュ値を書き込み、RecalcResult を返す。エンジンが無ければ None。"""
    formulas = collect_formulas(path)
    if not formulas:
        return {"ok": True, "errors": []}
    if not engine_available():
        return None
    values = compute_values(path, formulas)
    inject_cached_values(path, values)
    errors = find_errors(formulas, values)
    return {"ok": not errors, "errors": errors[:500]}
