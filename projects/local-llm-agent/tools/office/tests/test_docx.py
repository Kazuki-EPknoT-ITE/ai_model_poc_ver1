"""docx.*（T-15, T-16, T-18）。編集しない部分の XML が変わらないことを確かめる。"""

from __future__ import annotations

import json
import zipfile
from pathlib import Path
from typing import Any

import docx
import pytest
from conftest import FakeSampler, Runner
from docx.shared import RGBColor
from lxml import etree
from office_fixtures import make_docx, make_dotx

from office_tools.fsutil import sha256_file

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@pytest.fixture
def doc(ws: Path) -> Path:
    return make_docx(ws / "契約書.docx")


def c14n(el: Any) -> bytes:
    return bytes(etree.tostring(el, method="c14n"))


def body_children(path: Path) -> list[bytes]:
    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    body = root.find(f"{W}body")
    assert body is not None
    return [c14n(c) for c in body]


def other_parts(path: Path) -> dict[str, bytes]:
    """document.xml 以外の全パーツ（XML は正規化して比べる）。"""
    out = {}
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if name == "word/document.xml":
                continue
            data = z.read(name)
            if name.endswith((".xml", ".rels")):
                data = c14n(etree.fromstring(data))
            out[name] = data
    return out


def texts(run: Runner, path: str) -> dict[str, str]:
    env = run("docx.read", {"path": path, "include": ["headers"]})
    by_ref = {u["ref"]: u["text"] for u in env.get("untrusted", [])}
    return {b["id"]: by_ref.get(b.get("text_ref", ""), "") for b in env["data"]["blocks"]}


# ---------------------------------------------------------------- 読み取り


def test_read_blocks(run: Runner, doc: Path) -> None:
    env = run("docx.read", {"path": "契約書.docx"})
    blocks = env["data"]["blocks"]
    assert [b["id"] for b in blocks] == ["p:1", "p:2", "p:3", "p:4", "tbl:1", "p:5", "brk:1", "p:6", "p:7"]
    assert env["data"]["total_blocks"] == 9 and env["data"]["file_hash"] == sha256_file(doc)
    b = {x["id"]: x for x in blocks}
    assert b["p:1"]["kind"] == "heading" and b["p:1"]["level"] == 1 and b["p:1"]["style"] == "Heading 1"
    assert b["p:3"]["list"] == {"level": 0, "numbered": False}
    assert b["p:4"]["list"] == {"level": 0, "numbered": True}
    t = b["tbl:1"]
    assert t["rows"] == 3 and t["cols"] == 3
    assert t["cells"][0][1] == {"id": "tbl:1/r0c1", "merged": "h"}
    assert t["cells"][2][2] == {"id": "tbl:1/r2c2", "merged": "v"}
    unt = {u["ref"]: u["text"] for u in env["untrusted"]}
    assert unt["契約書.docx#p:2"] == "支払期限は30日以内とする。"
    assert "[r1c0] R1C0" in unt["契約書.docx#tbl:1"]
    assert "支払期限" not in json.dumps(env["data"], ensure_ascii=False)


def test_read_include_and_from_block(run: Runner, doc: Path) -> None:
    env = run(
        "docx.read",
        {"path": "契約書.docx", "include": ["headers", "images"], "from_block": "p:5", "max_blocks": 2},
    )
    assert [b["id"] for b in env["data"]["blocks"]] == ["p:5", "img:1"]
    assert env["continuation"]
    img = env["data"]["blocks"][1]
    assert img["width_mm"] == 20.0 and img["height_mm"] == 15.0
    env = run("docx.read", {"path": "契約書.docx", "include": ["headers"]})
    hdr = env["data"]["blocks"][0]
    assert hdr == {"id": "hdr:1", "kind": "header", "section": 1, "text_ref": "契約書.docx#hdr:1"}
    assert run.error("docx.read", {"path": "契約書.docx", "from_block": "p:99"})["code"] == "BLOCK_NOT_FOUND"


def test_outline(run: Runner, doc: Path) -> None:
    env = run("docx.outline", {"path": "契約書.docx"})
    d = env["data"]
    assert [(h["id"], h["level"]) for h in d["headings"]] == [("p:1", 1), ("p:6", 2)]
    assert d["tables"] == 1 and d["pages_estimate"] >= 1
    assert [u["text"] for u in env["untrusted"]] == ["第1条 目的", "第2条 金額"]


def test_unsupported_and_corrupt(run: Runner, ws: Path) -> None:
    (ws / "古い.doc").write_bytes(b"\xd0\xcf\x11\xe0")
    err = run.error("docx.read", {"path": "古い.doc"})
    assert err["code"] == "UNSUPPORTED_FORMAT" and "convert.to_docx" in err["message"]
    (ws / "壊れ.docx").write_bytes(b"PK\x03\x04garbage")
    assert run.error("docx.read", {"path": "壊れ.docx"})["code"] == "CORRUPT_FILE"


# ---------------------------------------------------------------- 編集


def test_replace_in_block_keeps_other_xml(run: Runner, doc: Path) -> None:
    before_body = body_children(doc)
    before_parts = other_parts(doc)
    env = run(
        "docx.edit",
        {
            "path": "契約書.docx",
            "ops": [{"op": "replace_in_block", "block_id": "p:2", "find": "30日以内", "replace": "45日以内"}],
        },
    )
    assert env["data"]["applied"] == 1 and env["data"]["results"] == [{"op_index": 0, "ok": True}]
    after_body = body_children(doc)
    assert len(after_body) == len(before_body)
    changed = [i for i, (a, b) in enumerate(zip(before_body, after_body, strict=True)) if a != b]
    assert changed == [1]  # p:2 だけが変わる
    assert other_parts(doc) == before_parts
    p = docx.Document(str(doc)).paragraphs[1]
    assert [(r.text, r.bold, r.italic) for r in p.runs] == [
        ("支払期限は", None, None),
        ("45日", True, None),
        ("以内とする。", None, True),
    ]


def test_replace_text_inherits_first_run_format(run: Runner, ws: Path) -> None:
    d = docx.Document()
    p = d.add_paragraph()
    r = p.add_run("赤字")
    r.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0, 0)
    p.add_run("普通")
    d.add_paragraph("そのまま")
    d.save(str(ws / "f.docx"))
    before = body_children(ws / "f.docx")
    run(
        "docx.edit",
        {"path": "f.docx", "ops": [{"op": "replace_text", "block_id": "p:1", "text": "新しい\n本文"}]},
    )
    p = docx.Document(str(ws / "f.docx")).paragraphs[0]
    assert len(p.runs) == 1 and p.runs[0].bold and p.runs[0].font.color.rgb == RGBColor(0xFF, 0, 0)
    assert p.text == "新しい\n本文"
    assert body_children(ws / "f.docx")[1:] == before[1:]


def test_ops_resolve_against_original_ids(run: Runner, doc: Path) -> None:
    run(
        "docx.edit",
        {
            "path": "契約書.docx",
            "ops": [
                {
                    "op": "insert_after",
                    "block_id": "p:1",
                    "blocks": [
                        {"kind": "paragraph", "text": "挿入A"},
                        {"kind": "heading", "level": 2, "text": "挿入B"},
                    ],
                },
                {"op": "insert_after", "block_id": "p:1", "blocks": [{"kind": "paragraph", "text": "挿入C"}]},
                {"op": "insert_before", "block_id": "p:7", "blocks": [{"kind": "page_break"}]},
                {"op": "delete_block", "block_id": "p:3"},
                {"op": "replace_text", "block_id": "p:4", "text": "手順X"},
                {"op": "delete_block", "block_id": "img:1"},
                {"op": "delete_block", "block_id": "brk:1"},
            ],
        },
    )
    t = texts(run, "契約書.docx")
    seq = [v for k, v in t.items() if k.startswith(("p:", "brk:"))]
    assert seq[:7] == [
        "第1条 目的",
        "挿入A",
        "挿入B",
        "挿入C",
        "支払期限は30日以内とする。",
        "手順X",
        "画像の前",
    ]
    env = run("docx.read", {"path": "契約書.docx", "include": ["images"]})
    kinds = [b["kind"] for b in env["data"]["blocks"]]
    assert "image" not in kinds and kinds.count("page_break") == 1


def test_table_ops(run: Runner, doc: Path) -> None:
    env = run(
        "docx.edit",
        {
            "path": "契約書.docx",
            "ops": [
                {"op": "table_set_cell", "cell_id": "tbl:1/r1c0", "text": "USD 12,500"},
                {
                    "op": "table_add_row",
                    "table_id": "tbl:1",
                    "after_row": 1,
                    "values": ["追加1", "追加2", "追加3"],
                },
                {"op": "table_add_row", "table_id": "tbl:1", "values": ["末尾"]},
                {"op": "table_delete_row", "table_id": "tbl:1", "row": 0},
                {"op": "table_set_cell", "cell_id": "tbl:1/r0c1", "text": "x"},
                {"op": "set_style", "block_id": "tbl:1", "style": "Light Grid Accent 1"},
                {"op": "set_style", "block_id": "p:7", "style": "Quote"},
            ],
        },
    )
    res = env["data"]["results"]
    assert [r["ok"] for r in res] == [True, True, True, True, False, True, True]
    assert "結合" in res[4]["error"]["message"]
    d = docx.Document(str(doc))
    t = d.tables[0]
    assert [c.text for c in t.rows[0].cells] == ["USD 12,500", "R1C1", "R1C2\nR2C2"]
    assert [c.text for c in t.rows[1].cells] == ["追加1", "追加2", "追加3"]
    assert t.rows[-1].cells[0].text == "末尾" and t.style.name == "Light Grid Accent 1"
    assert d.paragraphs[-1].style.name == "Quote"


def test_track_changes(run: Runner, doc: Path) -> None:
    run(
        "docx.edit",
        {
            "path": "契約書.docx",
            "track_changes": True,
            "author": "AI",
            "ops": [
                {"op": "replace_in_block", "block_id": "p:2", "find": "30日", "replace": "45日"},
                {"op": "replace_text", "block_id": "p:7", "text": "金額は USD 13,000 とする。"},
                {
                    "op": "insert_after",
                    "block_id": "p:7",
                    "blocks": [{"kind": "paragraph", "text": "第5条の2"}],
                },
                {"op": "delete_block", "block_id": "p:3"},
                {"op": "table_add_row", "table_id": "tbl:1", "values": ["a", "b", "c"]},
                {"op": "table_delete_row", "table_id": "tbl:1", "row": 1},
            ],
        },
    )
    with zipfile.ZipFile(doc) as z:
        xml = z.read("word/document.xml")
    root = etree.fromstring(xml)
    ins = root.findall(f".//{W}ins")
    dels = root.findall(f".//{W}del")
    assert ins and dels
    assert {e.get(f"{W}author") for e in ins + dels} == {"AI"}
    ids = [e.get(f"{W}id") for e in ins + dels]
    assert len(ids) == len(set(ids))
    deleted = "".join(t.text or "" for t in root.iter(f"{W}delText"))
    assert (
        "30" in deleted
        and "日" not in deleted.split("項目A")[0]
        and "項目A" in deleted
        and "金額は USD 12,500 とする。" in deleted
    )
    t = texts(run, "契約書.docx")
    assert t["p:2"] == "支払期限は45日以内とする。"
    assert any(v == "第5条の2" for v in t.values())
    assert root.find(f".//{W}trPr/{W}ins") is not None and root.find(f".//{W}trPr/{W}del") is not None
    # 変わる部分（30 → 45）だけを記録し、書式は太字の run を引き継ぐ
    ins_runs = [
        r
        for r in root.iter(f"{W}r")
        if r.getparent().tag == f"{W}ins" and "".join(t.text or "" for t in r.iter(f"{W}t")) == "45"
    ]
    assert ins_runs and ins_runs[0].find(f"{W}rPr/{W}b") is not None


def test_edit_errors_do_not_write(run: Runner, doc: Path) -> None:
    pre = sha256_file(doc)
    err = run.error(
        "docx.edit",
        {
            "path": "契約書.docx",
            "ops": [
                {"op": "replace_text", "block_id": "p:2", "text": "a"},
                {"op": "delete_block", "block_id": "p:57"},
            ],
        },
    )
    assert err["code"] == "BLOCK_NOT_FOUND" and "docx.read" in err["message"]
    err = run.error(
        "docx.edit",
        {
            "path": "契約書.docx",
            "expect_hash": "f" * 64,
            "ops": [{"op": "replace_text", "block_id": "p:2", "text": "a"}],
        },
    )
    assert err["code"] == "STALE_DOCUMENT"
    err = run.error(
        "docx.edit",
        {
            "path": "契約書.docx",
            "ops": [{"op": "replace_in_block", "block_id": "p:2", "find": "存在しない", "replace": "x"}],
        },
    )
    assert err["code"] == "INVALID_ARGS" and err["detail"]["results"][0]["ok"] is False
    assert (
        run.error("docx.edit", {"path": "契約書.docx", "ops": [{"op": "delete_block", "block_id": "hdr:1"}]})[
            "code"
        ]
        == "INVALID_ARGS"
    )
    assert sha256_file(doc) == pre


def test_partial_failure_and_header_edit(run: Runner, doc: Path) -> None:
    env = run(
        "docx.edit",
        {
            "path": "契約書.docx",
            "expect_hash": sha256_file(doc),
            "ops": [
                {"op": "replace_in_block", "block_id": "p:2", "find": "無い文字", "replace": "x"},
                {"op": "replace_in_block", "block_id": "hdr:1", "find": "社外秘", "replace": "部外秘"},
                {"op": "set_style", "block_id": "p:2", "style": "存在しないスタイル"},
            ],
        },
    )
    assert env["data"]["applied"] == 1 and env["warnings"]
    assert docx.Document(str(doc)).sections[0].header.paragraphs[0].text == "部外秘 ヘッダー"


def test_create_with_template(run: Runner, ws: Path) -> None:
    make_dotx(ws / "社内.dotx")
    env = run(
        "docx.create",
        {
            "path": "出力/報告書.docx",
            "template": "社内.dotx",
            "blocks": [
                {"kind": "heading", "level": 1, "text": "入渠報告"},
                {"kind": "paragraph", "text": "本書は…", "style": "社内本文"},
                {"kind": "paragraph", "text": "箇条", "list": {"level": 1, "numbered": False}},
                {"kind": "table", "header": ["項目", "金額"], "rows": [["塗装", "1,200,000"]]},
                {"kind": "page_break"},
                {"kind": "paragraph", "text": "x", "style": "無いスタイル"},
            ],
        },
    )
    assert env["data"]["artifact"]["op"] == "create"
    assert any("無いスタイル" in w for w in env["warnings"])
    d = docx.Document(str(ws / "出力" / "報告書.docx"))
    assert [p.style.name for p in d.paragraphs[:3]] == ["Heading 1", "社内本文", "List Bullet 2"]
    assert "雛形の本文" not in "\n".join(p.text for p in d.paragraphs)
    assert d.tables[0].cell(1, 1).text == "1,200,000" and d.tables[0].cell(0, 0).paragraphs[0].runs[0].bold
    blocks = run("docx.read", {"path": "出力/報告書.docx"})["data"]["blocks"]
    assert [b["kind"] for b in blocks] == [
        "heading",
        "paragraph",
        "paragraph",
        "table",
        "page_break",
        "paragraph",
    ]
    err = run.error("docx.create", {"path": "出力/報告書.docx", "blocks": [{"kind": "page_break"}]})
    assert err["code"] == "INVALID_ARGS"
    run("docx.create", {"path": "出力/報告書.docx", "blocks": [{"kind": "page_break"}], "overwrite": True})


def test_diff(run: Runner, doc: Path, ws: Path) -> None:
    import shutil

    shutil.copy(doc, ws / "v2.docx")
    run(
        "docx.edit",
        {
            "path": "v2.docx",
            "ops": [
                {"op": "replace_in_block", "block_id": "p:2", "find": "30日", "replace": "45日"},
                {"op": "delete_block", "block_id": "p:3"},
                {
                    "op": "insert_after",
                    "block_id": "p:7",
                    "blocks": [{"kind": "paragraph", "text": "追加条項"}],
                },
            ],
        },
    )
    env = run("docx.diff", {"a": "契約書.docx", "b": "v2.docx"})
    changes = env["data"]["changes"]
    unt = {u["ref"]: u["text"] for u in env["untrusted"]}
    kinds = [(c["kind"], c.get("a_id"), c.get("b_id")) for c in changes]
    assert (
        ("modify", "p:2", "p:2") in kinds
        and ("delete", "p:3", None) in kinds
        and ("insert", None, "p:7") in kinds
    )
    mod = next(c for c in changes if c["kind"] == "modify")
    assert "[-3-]{+4+}" in unt[mod["text_ref"]] or "[-30-]{+45+}" in unt[mod["text_ref"]]
    env = run("docx.diff", {"a": "契約書.docx", "b": "v2.docx", "granularity": "paragraph"})
    mod = next(c for c in env["data"]["changes"] if c["kind"] == "modify")
    assert {u["ref"]: u["text"] for u in env["untrusted"]}[mod["text_ref"]].startswith("- 支払期限は30日")


def test_comments(run: Runner, doc: Path) -> None:
    env = run("docx.comments", {"path": "契約書.docx"})
    assert env["data"]["comments"] == [] and "artifact" not in env["data"]
    env = run(
        "docx.comments",
        {
            "path": "契約書.docx",
            "add": [
                {"block_id": "p:2", "text": "金額を確認してください"},
                {"block_id": "tbl:1/r1c0", "text": "表の確認"},
            ],
        },
    )
    assert env["data"]["artifact"]["op"] == "overwrite"
    cs = env["data"]["comments"]
    assert [(c["block_id"], c["author"]) for c in cs] == [("p:2", "AI"), ("tbl:1", "AI")]
    assert {u["text"] for u in env["untrusted"]} == {"金額を確認してください", "表の確認"}
    env = run("docx.read", {"path": "契約書.docx", "include": ["comments"]})
    p2 = next(b for b in env["data"]["blocks"] if b["id"] == "p:2")
    assert p2["comments"][0]["author"] == "AI"


def test_images_extract_and_describe(run: Runner, doc: Path, ws: Path) -> None:
    env = run("docx.images", {"path": "契約書.docx"})
    img = env["data"]["images"][0]
    assert img["id"] == "img:1" and Path(img["path_extracted"]).read_bytes().startswith(b"\x89PNG")
    run.sampler = FakeSampler({"image_describe": "赤い四角形"})
    env = run("docx.images", {"path": "契約書.docx", "describe": True})
    assert env["untrusted"][0]["text"] == "赤い四角形" and env["untrusted"][0]["source"] == "ocr"
    assert run.sampler.requests[0].image_png is not None


def test_replace_in_block_insert_only_and_across_runs_tracked(run: Runner, ws: Path) -> None:
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("ABC")
    p.add_run("DEF").bold = True
    d.add_paragraph("未編集")
    d.save(str(ws / "g.docx"))
    before = body_children(ws / "g.docx")
    run(
        "docx.edit",
        {
            "path": "g.docx",
            "track_changes": True,
            "ops": [
                {"op": "replace_in_block", "block_id": "p:1", "find": "CD", "replace": "C-D"},
                {"op": "replace_in_block", "block_id": "p:1", "find": "AB", "replace": ""},
            ],
        },
    )
    assert texts(run, "g.docx")["p:1"] == "C-DEF"
    after = body_children(ws / "g.docx")
    assert after[1:] == before[1:]
    run(
        "docx.edit",
        {
            "path": "g.docx",
            "ops": [
                {"op": "replace_in_block", "block_id": "p:1", "find": "C-D", "replace": "XY", "count": 1}
            ],
        },
    )
    assert texts(run, "g.docx")["p:1"] == "XYEF"
