"""xlsx.*（T-11, T-12）。書式・数式・入力規則・マクロの保持を確かめる。"""

from __future__ import annotations

import datetime as dt
import zipfile
from pathlib import Path
from typing import Any

import pytest
from office_fixtures import VBA_BYTES, make_xlsm, make_xlsx
from office_testkit import Runner, needs_libreoffice
from openpyxl import Workbook, load_workbook

from office_tools.fsutil import sha256_file
from office_tools.xlsx import recalc as rc


@pytest.fixture
def book(ws: Path) -> Path:
    return make_xlsx(ws / "送金明細.xlsx")


def _snapshot(path: Path) -> dict[str, Any]:
    wb = load_workbook(path)
    ws = wb["明細"]
    return {
        "bold": ws["A1"].font.b,
        "fill": ws["A1"].fill.fgColor.rgb,
        "fmt_d2": ws["D2"].number_format,
        "fmt_a2": ws["A2"].number_format,
        "width_b": ws.column_dimensions["B"].width,
        "dv": [(str(d.sqref), d.formula1) for d in ws.data_validations.dataValidation],
        "cf": [str(cf.sqref) for cf in ws.conditional_formatting],
        "tables": {n: ws.tables[n].ref for n in ws.tables},
        "merged": [str(r) for r in ws.merged_cells.ranges],
        "f_sum": ws["D5"].value,
        "f_other": wb["集計"]["B1"].value,
    }


# ---------------------------------------------------------------- 読み取り


def test_info(run: Runner, book: Path) -> None:
    env = run("xlsx.info", {"path": "送金明細.xlsx"})
    d = env["data"]
    assert d["file_hash"] == sha256_file(book)
    assert [s["name"] for s in d["sheets"]] == ["明細", "集計"]
    s0 = d["sheets"][0]
    assert s0["merged"] == ["A7:C7"] and s0["tables"] == [{"name": "Meisai", "ref": "A1:D4"}]
    assert s0["used_range"] == "A1:D7" and d["has_macros"] is False


def _make_pivot_book(path: Path) -> None:
    """船名×月で金額を合計するピボットを持つブック（openpyxl で定義を組み立てる）。"""
    from openpyxl.pivot.cache import CacheDefinition, CacheField, CacheSource, SharedItems, WorksheetSource
    from openpyxl.pivot.record import RecordList
    from openpyxl.pivot.table import DataField, Location, PivotField, RowColField, TableDefinition

    wb = Workbook()
    data = wb.active
    assert data is not None
    data.title = "Data"
    for r in [["船名", "月", "金額"], ["平井丸", "9月", 100], ["平井丸", "10月", 200], ["第二丸", "9月", 50]]:
        data.append(r)
    pv = wb.create_sheet("集計")
    pv["A3"], pv["B3"] = "船名", "合計 / 金額"
    pv["A4"], pv["B4"] = "平井丸", 300
    pv["A5"], pv["B5"] = "第二丸", 50
    cache = CacheDefinition(
        cacheSource=CacheSource(type="worksheet", worksheetSource=WorksheetSource(ref="A1:C4", sheet="Data")),
        cacheFields=[CacheField(name=n, sharedItems=SharedItems()) for n in ("船名", "月", "金額")],
    )
    cache.records = RecordList()
    pt = TableDefinition(
        name="ピボット1",
        cacheId=1,
        dataCaption="値",
        location=Location(ref="A3:B5", firstHeaderRow=1, firstDataRow=1, firstDataCol=1),
        pivotFields=[
            PivotField(axis="axisRow", showAll=False),
            PivotField(axis="axisCol", showAll=False),
            PivotField(dataField=True, showAll=False),
        ],
        rowFields=[RowColField(x=0)],
        colFields=[RowColField(x=1)],
        dataFields=[DataField(name="合計 / 金額", fld=2, subtotal="sum")],
    )
    pt.cache = cache
    pv._pivots.append(pt)
    wb.save(path)


def test_info_reads_pivot_definitions(run: Runner, ws: Path) -> None:
    """XLS-09: ピボットは読み取りのみ。定義（場所・元データ・行・列・値）を返し、集計値は read_range で読む。"""
    _make_pivot_book(ws / "pivot.xlsx")
    sheets = {s["name"]: s for s in run("xlsx.info", {"path": "pivot.xlsx"})["data"]["sheets"]}
    assert sheets["Data"]["pivots"] == [] and sheets["Data"]["has_pivot"] is False
    (pivot,) = sheets["集計"]["pivots"]
    assert pivot == {
        "name": "ピボット1",
        "location": "A3:B5",
        "source": {"sheet": "Data", "ref": "A1:C4", "name": None},
        "rows": ["船名"],
        "columns": ["月"],
        "filters": [],
        "values": [{"name": "合計 / 金額", "field": "金額", "function": "sum"}],
    }
    cells = run("xlsx.read_range", {"path": "pivot.xlsx", "range": "集計!A3:B5"})["data"]
    assert "平井丸" in str(cells) and "300" in str(cells)


def test_read_range_values_formulas_formats(run: Runner, book: Path) -> None:
    env = run(
        "xlsx.read_range", {"path": "送金明細.xlsx", "range": "明細!A1:D5", "mode": "both", "header_row": 1}
    )
    d = env["data"]
    assert d["header"] == ["日付", "請求元", "区分", "金額"]
    assert d["rows"][0] == [{"date": "2026-10-01"}, "XX Shipping", "燃料", 12500.5]
    assert d["formulas"][3][3] == "=SUM(D2:D4)"
    assert d["formats"][0][3] == "#,##0.00"
    assert d["range"] == "'明細'!A1:D5"


def test_read_range_default_sheet_and_clamp(run: Runner, book: Path) -> None:
    env = run("xlsx.read_range", {"path": "送金明細.xlsx", "range": "A1:Z1000"})
    assert env["data"]["range"] == "'明細'!A1:D7"
    assert len(env["data"]["rows"]) == 7


def test_read_range_continuation(run: Runner, ws: Path) -> None:
    wb = Workbook()
    sh = wb.active
    assert sh is not None
    for i in range(1, 3001):
        sh.append([i, f"摘要{i}" * 3, i * 1.5])
    wb.save(ws / "big.xlsx")
    env = run("xlsx.read_range", {"path": "big.xlsx", "range": "Sheet!A1:C3000", "max_cells": 20000})
    rows = list(env["data"]["rows"])
    while "continuation" in env:
        env = run("xlsx.read_range", {"continuation": env["continuation"]})
        rows += env["data"]["rows"]
    assert len(rows) == 3000 and rows[-1][0] == 3000


def test_read_range_errors(run: Runner, book: Path, ws: Path) -> None:
    assert (
        run.error("xlsx.read_range", {"path": "送金明細.xlsx", "range": "無い!A1"})["code"] == "INVALID_ARGS"
    )
    assert (
        run.error("xlsx.read_range", {"path": "送金明細.xlsx", "range": "明細!ZZZZ1"})["code"]
        == "INVALID_ARGS"
    )
    (ws / "old.xls").write_bytes(b"\xd0\xcf\x11\xe0")
    err = run.error("xlsx.info", {"path": "old.xls"})
    assert err["code"] == "UNSUPPORTED_FORMAT" and "convert.to_xlsx" in err["message"]
    (ws / "broken.xlsx").write_bytes(b"not a zip")
    assert run.error("xlsx.info", {"path": "broken.xlsx"})["code"] == "CORRUPT_FILE"


def test_find(run: Runner, book: Path) -> None:
    env = run("xlsx.find", {"path": "送金明細.xlsx", "query": "燃料"})
    assert [h["cell"] for h in env["data"]["hits"]] == ["'明細'!C2", "'明細'!C4"]
    env = run(
        "xlsx.find", {"path": "送金明細.xlsx", "query": "shipping", "match": "contains", "sheet": "明細"}
    )
    assert env["data"]["hits"] == [{"cell": "'明細'!B2", "value": "XX Shipping"}]
    env = run("xlsx.find", {"path": "送金明細.xlsx", "query": "燃料", "max_results": 1})
    assert len(env["data"]["hits"]) == 1 and env["warnings"]


# ---------------------------------------------------------------- 書き込み


def test_write_cells_preserves_formatting(run: Runner, book: Path) -> None:
    before = _snapshot(book)
    pre = sha256_file(book)
    env = run(
        "xlsx.write_cells",
        {
            "path": "送金明細.xlsx",
            "sheet": "明細",
            "cells": [
                {"ref": "D3", "value": 3500},
                {"ref": "B8", "value": "=これは文字列"},
                {"ref": "A8", "value": {"date": "2026-10-11"}},
                {"ref": "D8", "formula": "=D5*2", "number_format": "#,##0"},
                {"ref": "明細!C8", "value": None},
            ],
            "expect_hash": pre,
        },
    )
    art = env["data"]["artifact"]
    assert art["pre_hash"] == pre and art["post_hash"] == sha256_file(book) and art["op"] == "overwrite"
    assert env["artifacts"] == [art]
    assert "recalc" not in env["data"]
    assert any("LibreOffice" in w for w in env["warnings"])
    assert _snapshot(book) == before
    wb = load_workbook(book)
    sh = wb["明細"]
    assert sh["D3"].value == 3500 and sh["D3"].number_format == "#,##0.00"
    assert sh["B8"].value == "=これは文字列" and sh["B8"].data_type == "s"
    assert sh["A8"].value == dt.datetime(2026, 10, 11) and sh["A8"].number_format == "yyyy-mm-dd"
    assert sh["D8"].value == "=D5*2"
    assert wb.calculation.fullCalcOnLoad is True


def test_write_cells_errors(run: Runner, book: Path) -> None:
    err = run.error(
        "xlsx.write_cells", {"path": "送金明細.xlsx", "sheet": "明細", "cells": [{"ref": "B7", "value": 1}]}
    )
    assert err["code"] == "INVALID_ARGS" and "結合" in err["message"]
    err = run.error(
        "xlsx.write_cells",
        {
            "path": "送金明細.xlsx",
            "sheet": "明細",
            "cells": [{"ref": "A1", "value": 1}],
            "expect_hash": "0" * 64,
        },
    )
    assert err["code"] == "STALE_DOCUMENT"
    # 写し間違いに気づけるよう、指定した値といまの file_hash の先頭を示す
    assert "先頭 00000000" in err["message"] and "expect_hash を省いて" in err["message"]
    err = run.error(
        "xlsx.write_cells", {"path": "送金明細.xlsx", "sheet": "無い", "cells": [{"ref": "A1", "value": 1}]}
    )
    assert err["code"] == "INVALID_ARGS"
    err = run.error(
        "xlsx.write_cells",
        {"path": "送金明細.xlsx", "sheet": "明細", "cells": [{"ref": "集計!A1", "value": 1}]},
    )
    assert err["code"] == "INVALID_ARGS"


def test_write_cells_file_locked(run: Runner, book: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import os

    def deny(src: object, dst: object) -> None:
        raise PermissionError("Excel で開いています")

    monkeypatch.setattr(os, "replace", deny)
    err = run.error(
        "xlsx.write_cells", {"path": "送金明細.xlsx", "sheet": "明細", "cells": [{"ref": "A1", "value": 1}]}
    )
    assert err["code"] == "FILE_LOCKED" and err["kind"] == "transient" and err["retry_after_ms"] > 0


def test_xlsm_keeps_vba(run: Runner, ws: Path) -> None:
    make_xlsm(ws / "マクロ.xlsm")
    env = run("xlsx.info", {"path": "マクロ.xlsm"})
    assert env["data"]["has_macros"] is True
    run(
        "xlsx.write_cells",
        {"path": "マクロ.xlsm", "sheet": "明細", "cells": [{"ref": "B2", "value": "変更"}]},
    )
    with zipfile.ZipFile(ws / "マクロ.xlsm") as z:
        assert z.read("xl/vbaProject.bin") == VBA_BYTES
        assert b"macroEnabled" in z.read("[Content_Types].xml")


def test_write_table(run: Runner, book: Path) -> None:
    env = run(
        "xlsx.write_table",
        {
            "path": "送金明細.xlsx",
            "sheet": "集計",
            "top_left": "D2",
            "header": ["日付", "金額"],
            "rows": [[{"date": "2026-10-01"}, 100], [{"date": "2026-10-02"}, 200.5]],
        },
    )
    assert env["data"]["range"] == "'集計'!D2:E4"
    sh = load_workbook(book)["集計"]
    assert sh["D2"].value == "日付" and sh["E4"].value == 200.5 and sh["D3"].value == dt.datetime(2026, 10, 1)


def test_insert_and_delete_rows_shift_references(run: Runner, book: Path) -> None:
    run("xlsx.insert_rows", {"path": "送金明細.xlsx", "sheet": "明細", "at": 3, "count": 2})
    wb = load_workbook(book)
    sh = wb["明細"]
    assert sh["D7"].value == "=SUM(D2:D6)"
    assert wb["集計"]["B1"].value == "=SUM('明細'!D2:D6)"
    assert wb["集計"]["B2"].value == "='明細'!D7*2"
    assert [str(r) for r in sh.merged_cells.ranges] == ["A9:C9"]
    assert sh.tables["Meisai"].ref == "A1:D6"
    assert str(sh.data_validations.dataValidation[0].sqref) == "C2:C12"
    assert [str(cf.sqref) for cf in sh.conditional_formatting] == ["D2:D6"]
    assert sh["B5"].value == "YY 商事" and sh["B5"].font.b is not True

    run("xlsx.delete_rows", {"path": "送金明細.xlsx", "sheet": "明細", "at": 3, "count": 2})
    assert _snapshot(book) == _snapshot(make_xlsx(book.with_name("orig.xlsx")))


def test_delete_rows_makes_ref_error(run: Runner, book: Path) -> None:
    run(
        "xlsx.write_cells",
        {"path": "送金明細.xlsx", "sheet": "集計", "cells": [{"ref": "C1", "formula": "='明細'!D3"}]},
    )
    run("xlsx.delete_rows", {"path": "送金明細.xlsx", "sheet": "明細", "at": 3, "count": 1})
    wb = load_workbook(book)
    assert wb["集計"]["C1"].value == "=#REF!"
    assert wb["明細"]["D4"].value == "=SUM(D2:D3)"


def test_insert_and_delete_cols(run: Runner, book: Path) -> None:
    run("xlsx.insert_cols", {"path": "送金明細.xlsx", "sheet": "明細", "at": "B", "count": 1})
    wb = load_workbook(book)
    sh = wb["明細"]
    assert sh["E5"].value == "=SUM(E2:E4)"
    assert wb["集計"]["B1"].value == "=SUM('明細'!E2:E4)"
    assert sh.tables["Meisai"].ref == "A1:E4"
    assert sh.column_dimensions["C"].width == 25
    assert [str(r) for r in sh.merged_cells.ranges] == ["A7:D7"]
    run("xlsx.delete_cols", {"path": "送金明細.xlsx", "sheet": "明細", "at": "B", "count": 1})
    sh = load_workbook(book)["明細"]
    assert sh["D5"].value == "=SUM(D2:D4)" and sh["B2"].value == "XX Shipping"


def test_add_and_rename_sheet(run: Runner, book: Path) -> None:
    run("xlsx.add_sheet", {"path": "送金明細.xlsx", "name": "メモ", "position": 0})
    assert load_workbook(book).sheetnames == ["メモ", "明細", "集計"]
    assert run.error("xlsx.add_sheet", {"path": "送金明細.xlsx", "name": "明細"})["code"] == "INVALID_ARGS"
    assert run.error("xlsx.add_sheet", {"path": "送金明細.xlsx", "name": "a/b"})["code"] == "INVALID_ARGS"
    run("xlsx.rename_sheet", {"path": "送金明細.xlsx", "name": "明細", "new_name": "送金 明細"})
    wb = load_workbook(book)
    assert wb["集計"]["B1"].value == "=SUM('送金 明細'!D2:D4)"
    assert "送金 明細" in wb.sheetnames


def test_create(run: Runner, ws: Path) -> None:
    env = run(
        "xlsx.create",
        {
            "path": "出力/送金明細.xlsx",
            "sheets": [
                {
                    "name": "明細",
                    "columns": [
                        {"header": "請求日", "width": 12, "number_format": "yyyy-mm-dd"},
                        {"header": "請求元", "width": 30},
                        {"header": "金額", "width": 14, "number_format": "#,##0.00"},
                    ],
                    "rows": [
                        [{"date": "2026-10-01"}, "XX Shipping", 12500.5],
                        [{"date": "2026-10-02"}, "YY", 10],
                    ],
                    "formulas": [{"ref": "C4", "formula": "=SUM(C2:C3)"}],
                    "freeze": "A2",
                    "as_table": {"name": "Remit"},
                }
            ],
        },
    )
    assert env["data"]["artifact"]["op"] == "create" and "pre_hash" not in env["data"]["artifact"]
    wb = load_workbook(ws / "出力" / "送金明細.xlsx")
    sh = wb["明細"]
    assert wb.sheetnames == ["明細"]
    assert sh["A1"].font.b and sh.freeze_panes == "A2" and sh.tables["Remit"].ref == "A1:C3"
    assert sh["C2"].number_format == "#,##0.00" and sh["C4"].value == "=SUM(C2:C3)"
    assert sh.column_dimensions["B"].width == 30
    err = run.error(
        "xlsx.create", {"path": "出力/送金明細.xlsx", "sheets": [{"name": "a", "columns": [{"header": "x"}]}]}
    )
    assert err["code"] == "INVALID_ARGS" and "overwrite" in err["message"]


def test_recalc_without_engine(run: Runner, book: Path) -> None:
    pre = sha256_file(book)
    env = run("xlsx.recalc", {"path": "送金明細.xlsx"})
    assert env["data"]["ok"] is False and env["data"]["recalculated"] is False
    assert env["warnings"] == [rc.NO_ENGINE_WARNING]
    assert sha256_file(book) == pre  # 書き換えない


def test_recalc_no_formulas_is_skipped(run: Runner, ws: Path) -> None:
    wb = Workbook()
    wb.active["A1"] = 1  # type: ignore[index]
    wb.save(ws / "plain.xlsx")
    env = run("xlsx.recalc", {"path": "plain.xlsx"})
    assert env["data"] == {"ok": True, "errors": []} and "artifacts" not in env


def test_recalc_with_fake_engine_injects_cached_values(
    run: Runner, book: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """LibreOffice の代わりに計算値を返し、キャッシュ値の書き込みとエラー一覧を確かめる。"""
    monkeypatch.setattr(rc, "engine_available", lambda: True)

    def fake_compute(path: Path, wanted: rc.Formulas) -> rc.Values:
        assert wanted == {
            "明細": {"D5": "=SUM(D2:D4)"},
            "集計": {"B1": "=SUM('明細'!D2:D4)", "B2": "='明細'!D5*2"},
        }
        return {"明細": {"D5": 15950.75}, "集計": {"B1": 15950.75, "B2": "#DIV/0!"}}

    monkeypatch.setattr(rc, "compute_values", fake_compute)
    before = _snapshot(book)
    env = run("xlsx.recalc", {"path": "送金明細.xlsx"})
    assert env["data"] == {
        "ok": False,
        "errors": [{"cell": "'集計'!B2", "value": "#DIV/0!", "formula": "='明細'!D5*2"}],
    }
    assert env["artifacts"][0]["post_hash"] == sha256_file(book)
    values = load_workbook(book, data_only=True)
    assert values["明細"]["D5"].value == 15950.75 and values["集計"]["B2"].value == "#DIV/0!"
    assert _snapshot(book) == before
    env = run(
        "xlsx.write_cells", {"path": "送金明細.xlsx", "sheet": "明細", "cells": [{"ref": "D2", "value": 1}]}
    )
    assert env["data"]["recalc"]["ok"] is False and "warnings" not in env


def test_chart(run: Runner, book: Path) -> None:
    run(
        "xlsx.chart",
        {
            "path": "送金明細.xlsx",
            "sheet": "集計",
            "kind": "bar",
            "data_range": "明細!B1:D4",
            "anchor": "D2",
            "title": "費用",
        },
    )
    wb = load_workbook(book)
    assert len(wb["集計"]._charts) == 1
    env = run("xlsx.info", {"path": "送金明細.xlsx"})
    assert env["data"]["sheets"][1]["has_chart"] is True
    run(
        "xlsx.chart",
        {
            "path": "送金明細.xlsx",
            "sheet": "集計",
            "kind": "scatter",
            "data_range": "明細!D1:D4",
            "anchor": "D20",
        },
        expect_ok=False,
    )


@needs_libreoffice
def test_recalc_with_libreoffice(run: Runner, book: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from office_tools import libreoffice

    monkeypatch.setattr(libreoffice, "find_soffice", lambda: __import__("conftest").REAL_SOFFICE)
    env = run("xlsx.recalc", {"path": "送金明細.xlsx"})
    assert env["data"]["ok"] is True
    assert load_workbook(book, data_only=True)["明細"]["D5"].value == pytest.approx(15950.75)


@pytest.mark.parametrize(
    ("formula", "expected"),
    [
        ("=SUM(A:A)", "=SUM(A:A)"),
        ("=SUM(2:3)", "=SUM(2:5)"),
        ("=SUM($A$1:$B$3)", "=SUM($A$1:$B$5)"),
        ("=A3+'Other S'!A3", "=A5+'Other S'!A3"),
        ("=明細!A3", "=明細!A5"),
        ("=MyName+1", "=MyName+1"),
        ("=Meisai[金額]", "=Meisai[金額]"),
        ('="A3"', '="A3"'),
    ],
)
def test_shift_formula_variants(formula: str, expected: str) -> None:
    from office_tools.xlsx.shift import Shift, ShiftRewriter, rewrite_formula

    assert rewrite_formula(formula, "明細", ShiftRewriter(Shift("明細", "row", 3, 2, False))) == expected


def test_shift_delete_and_cols() -> None:
    from office_tools.xlsx.shift import Shift, shift_body, shift_sqref

    d = Shift("S", "row", 3, 2, True)
    assert (
        shift_body("A3:B4", d) is None
        and shift_body("A2:B6", d) == "A2:B4"
        and shift_body("A4:B9", d) == "A3:B7"
    )
    assert shift_body("3:4", d) is None and shift_body("1:9", d) == "1:7"
    c = Shift("S", "col", 2, 1, True)
    assert shift_body("B:B", c) is None and shift_body("A:D", c) == "A:C" and shift_body("C5", c) == "B5"
    assert shift_body("$B$2", c) is None and shift_body("A1:C1", c) == "A1:B1"
    assert shift_sqref("A1:A3 B2 D4", c) == "A1:A3 C4"
    ins = Shift("S", "col", 2, 2, False)
    assert (
        shift_body("B:C", ins) == "D:E"
        and shift_body("A1", ins) == "A1"
        and shift_body("Name", ins) == "Name"
    )


def test_refs_parsing() -> None:
    from office_tools.errors import ToolFailure
    from office_tools.xlsx.refs import parse_range, quote_sheet, split_sheet

    assert split_sheet("'送金 明細'!A1") == ("送金 明細", "A1")
    assert split_sheet("'It''s'!B2") == ("It's", "B2")
    assert parse_range("'a b'!C3:A1").ref == "A1:C3"
    assert (
        quote_sheet("Sheet1") == "Sheet1" and quote_sheet("A1") == "'A1'" and quote_sheet("送金") == "'送金'"
    )
    with pytest.raises(ToolFailure):
        split_sheet("'abc")
