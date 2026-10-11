"""csv.* と data.transform（T-13）。"""

from __future__ import annotations

import datetime as dt
from decimal import Decimal
from pathlib import Path

import pytest
from conftest import Runner
from office_fixtures import make_xlsx
from openpyxl import load_workbook

from office_tools.errors import ToolFailure
from office_tools.tools.csvtool import detect_encoding
from office_tools.transform.expr import evaluate, parse

SJIS_CSV = '口座番号,名義,金額,日付\r\n0012345,ヤマダ　タロウ,"1,200",2026/10/01\r\n0099999,株式会社サンプル,3000,2026/10/02\r\n'


# ---------------------------------------------------------------- csv


def test_csv_read_shift_jis(run: Runner, ws: Path) -> None:
    (ws / "銀行明細.csv").write_bytes(SJIS_CSV.encode("cp932"))
    env = run("csv.read", {"path": "銀行明細.csv"})
    d = env["data"]
    assert d["encoding"] == "cp932" and d["delimiter"] == ","
    assert d["header"] == ["口座番号", "名義", "金額", "日付"]
    assert d["rows"][0] == ["0012345", "ヤマダ　タロウ", "1,200", "2026/10/01"]
    assert d["total_rows"] == 2


def test_csv_read_utf8_bom_and_tab(run: Runner, ws: Path) -> None:
    (ws / "a.csv").write_bytes("﻿品名\t数量\nりんご\t3\n".encode())
    env = run("csv.read", {"path": "a.csv"})
    assert env["data"]["encoding"] == "utf-8-sig" and env["data"]["delimiter"] == "\t"
    assert env["data"]["header"] == ["品名", "数量"]


def test_csv_read_explicit_encoding_and_offset(run: Runner, ws: Path) -> None:
    (ws / "a.csv").write_bytes(SJIS_CSV.encode("cp932"))
    env = run("csv.read", {"path": "a.csv", "encoding": "cp932", "offset": 1, "max_rows": 1})
    assert env["data"]["rows"] == [["0099999", "株式会社サンプル", "3000", "2026/10/02"]]
    env = run("csv.read", {"path": "a.csv", "encoding": "utf-8"})
    assert env["warnings"]


def test_detect_encoding_variants() -> None:
    text = "請求番号,取引先,摘要\n" + "INV-1,東京海運株式会社,入渠費用の支払い\n" * 20
    assert detect_encoding(text.encode("cp932")) == "cp932"
    assert detect_encoding(text.encode("utf-8")) == "utf-8"
    assert detect_encoding(b"\xef\xbb\xbf" + text.encode("utf-8")) == "utf-8-sig"


def test_csv_continuation(run: Runner, ws: Path) -> None:
    lines = ["no,摘要"] + [f"{i},{'テキスト' * 10}" for i in range(5000)]
    (ws / "big.csv").write_text("\n".join(lines), encoding="utf-8")
    env = run("csv.read", {"path": "big.csv", "max_rows": 20000})
    rows = list(env["data"]["rows"])
    assert env["continuation"]
    while "continuation" in env:
        env = run("csv.read", {"continuation": env["continuation"]})
        rows += env["data"]["rows"]
    assert len(rows) == 5000 and rows[-1][0] == "4999"


def test_csv_write_cp932_and_errors(run: Runner, ws: Path) -> None:
    env = run(
        "csv.write",
        {
            "path": "出力/振込.csv",
            "header": ["口座", "金額", "日付"],
            "rows": [["0012345", 1000, {"date": "2026-10-01"}], ["1", 0.1, True]],
            "encoding": "cp932",
        },
    )
    assert env["data"]["artifact"]["op"] == "create"
    text = (ws / "出力" / "振込.csv").read_bytes().decode("cp932")
    assert text == "口座,金額,日付\r\n0012345,1000,2026-10-01\r\n1,0.1,TRUE\r\n"
    err = run.error("csv.write", {"path": "出力/振込.csv", "rows": [["a"]]})
    assert err["code"] == "INVALID_ARGS"
    err = run.error("csv.write", {"path": "x.csv", "rows": [["🚢"]], "encoding": "cp932"})
    assert err["code"] == "INVALID_ARGS" and "utf-8-sig" in err["message"]
    run("csv.write", {"path": "y.csv", "rows": [["🚢"]]})
    assert (ws / "y.csv").read_bytes().startswith(b"\xef\xbb\xbf")


# ---------------------------------------------------------------- 式


def ev(src: str, **row: object) -> object:
    return evaluate(parse(src), row)


def test_expr_arithmetic_is_decimal() -> None:
    assert ev("a + b", a=Decimal("0.1"), b=Decimal("0.2")) == Decimal("0.3")
    assert ev("1 + 2 * 3 - 4 / 2") == Decimal("5")
    assert ev("-(1 + 2) * 2") == Decimal("-6")
    assert ev("round(10 / 3, 2)") == Decimal("3.33")
    assert ev("round(2.5)") == Decimal("3")
    assert ev("a / 0", a=Decimal(1)) is None
    assert ev("a + 1", a=None) is None


def test_expr_logic_and_functions() -> None:
    row = {"金額": Decimal(12500), "請求元": "XX Shipping", "日付": dt.date(2026, 10, 1), "備考": None}
    assert evaluate(parse("金額 >= 10000 and contains(請求元, 'Shipping')"), row) is True
    assert evaluate(parse("not startswith(請求元, 'YY') or 金額 < 0"), row) is True
    assert evaluate(parse("日付 >= date('2026-09-30')"), row) is True
    assert evaluate(parse("日付 == '2026-10-01'"), row) is True
    assert evaluate(parse("coalesce(備考, 'なし')"), row) == "なし"
    assert evaluate(parse("備考 == null"), row) is True
    assert evaluate(parse('`請求元` = "XX Shipping"'), row) is True
    assert evaluate(parse("[金額] != 1"), row) is True
    assert evaluate(parse("'a' + 'b'"), row) == "ab"
    assert evaluate(parse("日付 + 30"), row) == dt.date(2026, 10, 31)


@pytest.mark.parametrize(
    "src",
    ["__import__('os').system('x')", "a.b", "1 +", "(1", "a = = 1", "open('x')", "a; b", "lambda: 1", "x[0]"],
)
def test_expr_rejects_non_grammar(src: str) -> None:
    with pytest.raises(ToolFailure) as e:
        parse(src)
    assert e.value.code == "INVALID_ARGS"


def test_expr_type_errors() -> None:
    with pytest.raises(ToolFailure):
        ev("a * 2", a="abc")
    with pytest.raises(ToolFailure):
        ev("a < 1", a=True)


# ---------------------------------------------------------------- data.transform


def test_transform_reconcile_xlsx_and_csv(run: Runner, ws: Path) -> None:
    make_xlsx(ws / "送金明細.xlsx")
    (ws / "銀行.csv").write_bytes(
        '請求元,金額\r\nXX Shipping,"12,500.50"\r\nYY 商事,2999\r\nQQ,1\r\n'.encode("cp932")
    )
    env = run(
        "data.transform",
        {
            "inputs": [
                {"name": "inv", "source": {"xlsx": "送金明細.xlsx", "range": "明細!A1:D4", "header_row": 1}},
                {"name": "bank", "source": {"csv": "銀行.csv"}},
            ],
            "steps": [
                {"op": "reconcile", "with": "bank", "key": ["請求元"], "compare": ["金額"], "tolerance": 0},
                {"op": "filter", "where": "_status != 'match'"},
                {"op": "select", "columns": ["請求元", "金額", "金額_bank", "_status"]},
            ],
            "output": {"xlsx": "突合結果.xlsx", "sheet": "差異", "as_table": True},
        },
    )
    d = env["data"]
    assert d["columns"] == ["請求元", "金額", "金額_bank", "_status"]
    assert d["preview"] == [
        ["YY 商事", 3000, 2999, "mismatch"],
        ["ZZ Marine", 450.25, None, "left_only"],
        ["QQ", None, 1, "right_only"],
    ]
    assert d["rows_out"] == 3 and d["artifact"]["op"] == "create"
    sh = load_workbook(ws / "突合結果.xlsx")["差異"]
    assert sh["A1"].value == "請求元" and sh["D2"].value == "mismatch" and len(sh.tables) == 1


def test_transform_tolerance(run: Runner) -> None:
    env = run(
        "data.transform",
        {
            "inputs": [
                {"name": "a", "source": {"header": ["k", "v"], "rows": [["x", 100.0], ["y", 50]]}},
                {"name": "b", "source": {"header": ["k", "v"], "rows": [["x", 100.4], ["y", 51]]}},
            ],
            "steps": [{"op": "reconcile", "with": "b", "key": ["k"], "compare": ["v"], "tolerance": 0.5}],
            "output": {"return_only": True},
        },
    )
    assert [r[-2] for r in env["data"]["preview"]] == ["match", "mismatch"]
    assert "artifact" not in env["data"]


def test_transform_group_sort_derive_rename_dedupe_join(run: Runner, ws: Path) -> None:
    rows = [
        ["燃料", "XX", 0.1, {"date": "2026-10-01"}],
        ["燃料", "XX", 0.2, {"date": "2026-10-02"}],
        ["港費", "YY", 300, {"date": "2026-10-02"}],
        ["港費", "YY", 300, {"date": "2026-10-02"}],
        ["その他", None, None, None],
    ]
    env = run(
        "data.transform",
        {
            "inputs": [
                {"name": "m", "source": {"header": ["区分", "取引先", "金額", "日付"], "rows": rows}},
                {
                    "name": "t",
                    "source": {
                        "header": ["取引先", "名称"],
                        "rows": [["XX", "XX Shipping"], ["ZZ", "ZZ Marine"]],
                    },
                },
            ],
            "steps": [
                {"op": "dedupe", "by": ["区分", "取引先", "金額", "日付"]},
                {"op": "derive", "column": "税込", "expr": "round(金額 * 1.1, 2)"},
                {"op": "join", "with": "t", "on": [{"left": "取引先", "right": "取引先"}], "how": "left"},
                {
                    "op": "group",
                    "by": ["区分"],
                    "aggs": [
                        {"column": "金額", "fn": "sum", "as": "合計"},
                        {"column": "金額", "fn": "count", "as": "件数"},
                        {"column": "金額", "fn": "avg", "as": "平均"},
                        {"column": "日付", "fn": "max", "as": "最終日"},
                        {"column": "名称", "fn": "first", "as": "名称"},
                    ],
                },
                {"op": "sort", "by": [{"column": "合計", "desc": True}]},
                {"op": "rename", "map": {"区分": "費目"}},
            ],
            "output": {"csv": "集計.csv"},
        },
    )
    assert env["data"]["columns"] == ["費目", "合計", "件数", "平均", "最終日", "名称"]
    assert env["data"]["preview"] == [
        ["港費", 300, 1, 300, {"date": "2026-10-02"}, None],
        ["燃料", 0.3, 2, 0.15, {"date": "2026-10-02"}, "XX Shipping"],
        ["その他", 0, 0, None, None, None],
    ]
    text = (ws / "集計.csv").read_bytes().decode("utf-8-sig")
    assert text.splitlines()[2] == "燃料,0.3,2,0.15,2026-10-02,XX Shipping"


def test_transform_outer_join_and_errors(run: Runner) -> None:
    base = {"name": "a", "source": {"header": ["k", "v"], "rows": [["1", 1], ["2", 2]]}}
    other = {"name": "b", "source": {"header": ["k", "v"], "rows": [["2", 20], ["3", 30]]}}
    env = run(
        "data.transform",
        {
            "inputs": [base, other],
            "steps": [{"op": "join", "with": "b", "on": [{"left": "k", "right": "k"}], "how": "outer"}],
            "output": {"return_only": True},
        },
    )
    assert env["data"]["columns"] == ["k", "v", "v_b"]
    assert env["data"]["preview"] == [["1", 1, None], ["2", 2, 20], ["3", None, 30]]
    err = run.error(
        "data.transform",
        {
            "inputs": [base],
            "steps": [{"op": "filter", "where": "無い列 > 1"}],
            "output": {"return_only": True},
        },
    )
    assert err["code"] == "INVALID_ARGS" and "steps[0]" in err["message"]
    err = run.error(
        "data.transform",
        {
            "inputs": [base],
            "steps": [{"op": "join", "with": "zz", "on": [{"left": "k", "right": "k"}], "how": "inner"}],
            "output": {"return_only": True},
        },
    )
    assert err["code"] == "INVALID_ARGS"
    err = run.error(
        "data.transform",
        {
            "inputs": [{"name": "r", "source": {"result_of": "batch-1"}}],
            "steps": [],
            "output": {"return_only": True},
        },
    )
    assert err["code"] == "INVALID_ARGS" and "Runtime" in err["message"]
