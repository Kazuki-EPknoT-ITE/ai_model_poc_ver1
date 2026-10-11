"""テスト用の文書を合成する（実データは使わない）。"""

from __future__ import annotations

import datetime as dt
import io
import zipfile
from pathlib import Path

import docx
from docx.enum.text import WD_BREAK
from docx.shared import Mm
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas


def png_bytes(color: str = "red", size: tuple[int, int] = (40, 30)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, format="PNG")
    return buf.getvalue()


# ---------------------------------------------------------------- Word


def make_docx(path: Path) -> Path:
    """見出し・書式付きの段落・箇条書き・結合セルのある表・画像・改ページ・ヘッダー/フッター。

    ブロックID: p:1 見出し1 / p:2 書式付き段落 / p:3 箇条書き / p:4 番号付き / tbl:1 /
    p:5 / img:1 / brk:1 / p:6 見出し2 / p:7 / hdr:1 / ftr:1
    """
    d = docx.Document()
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "社外秘 ヘッダー"
    sec.footer.paragraphs[0].text = "フッター 1"
    d.add_heading("第1条 目的", level=1)
    p = d.add_paragraph()
    p.add_run("支払期限は")
    r = p.add_run("30日")
    r.bold = True
    r = p.add_run("以内とする。")
    r.italic = True
    d.add_paragraph("項目A", style="List Bullet")
    d.add_paragraph("手順1", style="List Number")
    t = d.add_table(rows=3, cols=3)
    t.style = "Table Grid"
    for i in range(3):
        for j in range(3):
            t.cell(i, j).text = f"R{i}C{j}"
    t.cell(0, 0).merge(t.cell(0, 1))
    t.cell(0, 0).paragraphs[0].text = "見出しセル"
    for extra in t.cell(0, 0).paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    t.cell(1, 2).merge(t.cell(2, 2))
    d.add_paragraph("画像の前")
    d.add_picture(io.BytesIO(png_bytes()), width=Mm(20))
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    d.add_heading("第2条 金額", level=2)
    d.add_paragraph("金額は USD 12,500 とする。")
    d.save(str(path))
    return path


def make_dotx(path: Path) -> Path:
    """テンプレート（.dotx）: 独自の段落スタイルを持つ。"""
    d = docx.Document()
    from docx.enum.style import WD_STYLE_TYPE

    st = d.styles.add_style("社内本文", WD_STYLE_TYPE.PARAGRAPH)
    st.font.size = docx.shared.Pt(10.5)
    d.add_paragraph("雛形の本文（消える）")
    tmp = path.with_suffix(".docx.tmp")
    d.save(str(tmp))
    with zipfile.ZipFile(tmp) as src, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "[Content_Types].xml":
                data = data.replace(
                    b"wordprocessingml.document.main+xml", b"wordprocessingml.template.main+xml"
                )
            out.writestr(info, data)
    tmp.unlink()
    return path


# ---------------------------------------------------------------- Excel


def make_xlsx(path: Path) -> Path:
    """書式・数式・入力規則・条件付き書式・結合セル・テーブル・2シート間の参照。"""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "明細"
    header = ["日付", "請求元", "区分", "金額"]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FFFF00")
    rows = [
        (dt.date(2026, 10, 1), "XX Shipping", "燃料", 12500.5),
        (dt.date(2026, 10, 2), "YY 商事", "港費", 3000),
        (dt.date(2026, 10, 3), "ZZ Marine", "燃料", 450.25),
    ]
    for row in rows:
        ws.append(list(row))
    for r in range(2, 5):
        ws.cell(row=r, column=1).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=4).number_format = "#,##0.00"
    ws["D5"] = "=SUM(D2:D4)"
    ws["D5"].number_format = "#,##0.00"
    ws["A5"] = "合計"
    ws.merge_cells("A7:C7")
    ws["A7"] = "備考（結合セル）"
    ws.column_dimensions["B"].width = 25
    dv = DataValidation(type="list", formula1='"燃料,港費,その他"', allow_blank=True)
    dv.add("C2:C10")
    ws.add_data_validation(dv)
    ws.conditional_formatting.add(
        "D2:D4",
        CellIsRule(operator="greaterThan", formula=["10000"], fill=PatternFill("solid", fgColor="FFC7CE")),
    )
    table = Table(displayName="Meisai", ref="A1:D4")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(table)
    s2 = wb.create_sheet("集計")
    s2["A1"] = "総額"
    s2["B1"] = "=SUM('明細'!D2:D4)"
    s2["B2"] = "='明細'!D5*2"
    wb.save(path)
    return path


def make_xlsm(path: Path) -> Path:
    """VBA プロジェクト（ダミー）を含む .xlsm。"""
    tmp = path.with_suffix(".xlsx")
    make_xlsx(tmp)
    with zipfile.ZipFile(tmp) as src, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                ).replace(
                    b'<Default Extension="xml"',
                    b'<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/><Default Extension="xml"',
                )
            if info.filename == "xl/_rels/workbook.xml.rels":
                data = data.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdVba" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/></Relationships>',
                )
            out.writestr(info, data)
        out.writestr("xl/vbaProject.bin", VBA_BYTES)
    tmp.unlink()
    return path


VBA_BYTES = b"\xd0\xcf\x11\xe0dummy-vba-project-for-test" * 4


# ---------------------------------------------------------------- PDF


def make_pdf(path: Path, *, encrypt: str | None = None) -> Path:
    """1: 本文 / 2: 2段組 / 3: 罫線の表 / 4: 画像だけ（スキャン相当）。"""
    c = canvas.Canvas(str(path), pagesize=A4, encrypt=encrypt)
    c.setTitle("Test Document")
    w, h = A4
    c.setFont("Helvetica", 12)
    y = h - 72
    for line in ["Charter Party Summary", "Owner: Example Shipping Co.", "Hire: USD 12,500 per day"]:
        c.drawString(72, y, line)
        y -= 18
    c.showPage()
    c.setFont("Helvetica", 11)
    y = h - 72
    for i in range(1, 13):
        c.drawString(60, y, f"Left column line {i}")
        c.drawString(330, y, f"Right column line {i}")
        y -= 16
    c.showPage()
    c.setFont("Helvetica", 10)
    x0, y0, cw, rh = 72, h - 120, 120, 20
    data = [["Item", "Qty", "Amount"], ["Fuel", "2", "1000"], ["Port", "1", "300"]]
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            c.rect(x0 + j * cw, y0 - i * rh, cw, rh)
            c.drawString(x0 + j * cw + 4, y0 - i * rh + 6, val)
    c.showPage()
    img = ImageReader(io.BytesIO(png_bytes("gray", (200, 280))))
    c.drawImage(img, 72, h - 400, width=200, height=280)
    c.showPage()
    c.save()
    return path


def make_invoice_pdf(path: Path, *, scanned_second_page: bool = False) -> Path:
    c = canvas.Canvas(str(path), pagesize=A4)
    w, h = A4
    c.setFont("Helvetica", 11)
    lines = [
        "INVOICE",
        "Issuer: Example Shipping Co., Ltd.",
        "Registration No: T1234567890123",
        "Bill To: Sample Marine K.K.",
        "Vessel: MV EXAMPLE STAR  Voyage: 2610E",
        "Issue Date: 2026-10-01   Due Date: 2026-10-31",
        "Currency: USD",
        "Bunker fuel   2 x 5,000.00 = 10,000.00",
        "Port charges  1 x 2,500.00 = 2,500.00",
        "Subtotal: 12,500.00",
        "Tax: 0.00",
        "Total: 12,500.00",
        "Bank: Example Bank  Branch: Tokyo  Account: Ordinary 1234567  SWIFT: EXAMJPJT",
    ]
    y = h - 72
    for line in lines:
        c.drawString(72, y, line)
        y -= 18
    c.showPage()
    if scanned_second_page:
        img = ImageReader(io.BytesIO(png_bytes("white", (300, 400))))
        c.drawImage(img, 72, h - 500, width=300, height=400)
        c.showPage()
    c.save()
    return path
