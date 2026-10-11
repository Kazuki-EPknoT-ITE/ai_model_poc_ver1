"""試験用の部品（以前は conftest にあった。全体で pytest を回すと conftest の名前が衝突するため分けた）。"""

from __future__ import annotations

import math
import re
from collections.abc import Callable
from pathlib import Path
from typing import Any

import docx
import pytest
from openpyxl import Workbook
from qdrant_client import QdrantClient
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

from llm_rag.analyzer import BigramTokenizer
from llm_rag.config import AclMap, AclRule, ChunkSettings, EmbedderSettings, RagSettings
from llm_rag.index import RagIndex
from llm_rag.search import Searcher

JP_FONT = "HeiseiKakuGo-W5"
pdfmetrics.registerFont(UnicodeCIDFont(JP_FONT))


# ---------------------------------------------------------------- 文書の合成


def write_docx(path: Path, blocks: list[tuple[str, Any]]) -> Path:
    """blocks: ("h1"|"h2", text) / ("p", text) / ("table", rows)。"""
    d = docx.Document()
    for kind, val in blocks:
        if kind.startswith("h"):
            d.add_heading(str(val), level=int(kind[1:]))
        elif kind == "p":
            d.add_paragraph(str(val))
        elif kind == "table":
            rows: list[list[str]] = val
            t = d.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, v in enumerate(row):
                    t.cell(r, c).text = v
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))
    return path


def write_pdf(path: Path, pages: list[list[str]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(path), pagesize=A4)
    _, h = A4
    for lines in pages:
        c.setFont(JP_FONT, 11)
        y = h - 72
        for ln in lines:
            c.drawString(60, y, ln)
            y -= 16
        c.showPage()
    c.save()
    return path


def write_xlsx(path: Path, sheets: dict[str, list[list[Any]]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        assert ws is not None
        ws.title = name
        first = False
        for row in rows:
            ws.append(row)
    wb.save(str(path))
    return path


def write_text(path: Path, text: str, encoding: str = "utf-8") -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(text.encode(encoding))
    return path


# ---------------------------------------------------------------- 試験用の埋め込み


class ConceptEmbedder:
    """小さな同義語辞書で「意味の近さ」を模した埋め込み（ハイブリッド検索の試験用）。

    同じ概念の語（支払・振込・送金 など）は同じ次元に数える。語が無い文は "other" 次元だけを持つ。
    """

    CONCEPTS: dict[str, list[str]] = {
        "pay": ["支払", "振込", "送金", "入金"],
        "due": ["期日", "期限", "締切", "締め切り"],
        "meeting": ["会議", "打合せ", "ミーティング"],
        "room": ["会議室", "部屋"],
        "leave": ["休暇", "有給", "休み"],
        "ship": ["船舶", "本船", "傭船"],
    }

    def __init__(self) -> None:
        self._keys = sorted(self.CONCEPTS)

    @property
    def model_id(self) -> str:
        return "concept-test"

    @property
    def dim(self) -> int:
        return len(self._keys) + 1

    def _vec(self, text: str) -> list[float]:
        v = [0.0] * self.dim
        for i, k in enumerate(self._keys):
            v[i] = float(sum(len(re.findall(w, text)) for w in self.CONCEPTS[k]))
        v[-1] = 0.3
        n = math.sqrt(sum(x * x for x in v))
        return [x / n for x in v]

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [self._vec(t) for t in texts]

    def embed_query(self, text: str) -> list[float]:
        return self._vec(text)


# ---------------------------------------------------------------- 設定・索引


def make_settings(
    tmp_path: Path,
    rules: list[tuple[str, list[str]] | tuple[str, list[str], str]],
    *,
    sources: list[Path] | None = None,
    max_tokens: int = 800,
    overlap_tokens: int = 100,
    **extra: Any,
) -> RagSettings:
    acl = AclMap(
        rules=[
            AclRule(prefix=r[0], acl_groups=r[1], doc_type=(r[2] if len(r) > 2 else None))  # type: ignore[misc]
            for r in rules
        ]
    )
    return RagSettings(
        org_id="test",
        sources=[str(p) for p in (sources or [tmp_path / "docs"])],
        index_dir=str(tmp_path / "index"),
        acl_map=acl,
        tokenizer="bigram",
        embedder=EmbedderSettings(kind="hash", dim=128),
        chunking=ChunkSettings(max_tokens=max_tokens, overlap_tokens=overlap_tokens),
        **extra,
    )


IndexFactory = Callable[..., RagIndex]


@pytest.fixture
def make_index() -> IndexFactory:
    def factory(settings: RagSettings, embedder: Any = None, client: QdrantClient | None = None) -> RagIndex:
        return RagIndex(
            settings,
            client=client or QdrantClient(location=":memory:"),
            embedder=embedder,
            tokenizer=BigramTokenizer(),
        )

    return factory


def paths_of(searcher: Searcher, query: str, groups: list[str], **kw: Any) -> list[str]:
    return [Path(h.path).name for h in searcher.search(query, groups, **kw)]
