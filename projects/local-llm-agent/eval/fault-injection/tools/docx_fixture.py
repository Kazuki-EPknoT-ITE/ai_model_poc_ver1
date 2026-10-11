"""FI-17 の道具: 試験用の Word 文書を作る / 他人の編集を模擬して段落を足す。

python docx_fixture.py make <path>          決定的な docx を作る（ハッシュをシナリオに書くため）
python docx_fixture.py append <path> <text> 段落を1つ足して保存する（他人による同時編集）
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document


def make(path: Path) -> None:
    from lla_eval.synth import stabilize_zip  # noqa: PLC0415 - PYTHONPATH に eval/harness を入れて呼ぶ

    doc = Document()
    doc.add_heading("入渠報告書", level=1)
    doc.add_paragraph("本文: 旧")
    doc.add_paragraph("担当: 未定")
    doc.save(str(path))
    stabilize_zip(path)


def append(path: Path, text: str) -> None:
    doc = Document(str(path))
    doc.add_paragraph(text)
    doc.save(str(path))


if __name__ == "__main__":
    cmd, target = sys.argv[1], Path(sys.argv[2])
    if cmd == "make":
        make(target)
    elif cmd == "append":
        append(target, sys.argv[3])
    else:
        raise SystemExit(f"unknown command: {cmd}")
