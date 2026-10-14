"""採点（local_llm_test_spec.md §3.2）。

- invoice_fields: 帳票抽出。項目ごとに完全一致（日付は ISO、金額は Decimal、文字列は NFKC＋空白除去）。
  金額・日付とその他を分けて集計し、needs_review（確信度が低い・検算に失敗）の付いた誤りは
  「検知できた誤り」として別に数える。
- artifact_checks: ツール操作。期待成果物との比較（xlsx はセル値、docx は段落テキスト、csv は行）など。
- doc_quality: 文書品質。壊れていないか（openpyxl / python-docx で開けるか、zip が正しいか）と、
  編集していない要素が変わっていないか（書式崩れ）。
- resilience: 継続実行。障害注入（eval/fault-injection）の結果から完了率と重複書き込み件数を出す。
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

REVIEW_THRESHOLD = 0.8  # office-tools の invoice.extract と同じ（これ未満は needs_review 扱い）
DATE_FIELDS = {"issue_date", "due_date"}
AMOUNT_FIELDS = {"subtotal", "tax", "total"}
AMOUNT_LINE_FIELDS = {"amount", "unit_price"}
NUMBER_LINE_FIELDS = {"qty"}
_MONTHS = {
    m: i
    for i, ms in enumerate(
        (
            ("jan", "january"),
            ("feb", "february"),
            ("mar", "march"),
            ("apr", "april"),
            ("may",),
            ("jun", "june"),
            ("jul", "july"),
            ("aug", "august"),
            ("sep", "sept", "september"),
            ("oct", "october"),
            ("nov", "november"),
            ("dec", "december"),
        ),
        start=1,
    )
    for m in ms
}


# ---------------------------------------------------------------- 正規化
def norm_text(v: Any) -> str | None:
    if v is None:
        return None
    s = unicodedata.normalize("NFKC", str(v))
    s = re.sub(r"\s+", "", s)
    return s or None


def norm_amount(v: Any) -> Decimal | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, int | float):
        return Decimal(repr(v)) if isinstance(v, float) else Decimal(v)
    s = unicodedata.normalize("NFKC", str(v)).strip()
    neg = s.startswith("(") and s.endswith(")")
    s = re.sub(r"[^\d.\-]", "", s.replace(",", ""))
    if not s or s in ("-", ".", "-."):
        return None
    try:
        d = Decimal(s)
    except InvalidOperation:
        return None
    return -d if neg else d


def norm_date(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, dict) and "date" in v:
        v = v["date"]
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = unicodedata.normalize("NFKC", str(v)).strip()
    m = re.fullmatch(r"(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})日?(?:[ T].*)?", s)
    if m:
        y, mo, d = (int(x) for x in m.groups())
    else:
        m = re.fullmatch(r"(\d{1,2})[\s-]+([A-Za-z]{3,9})\.?[\s,-]+(\d{4})", s)
        m2 = re.fullmatch(r"([A-Za-z]{3,9})\.?\s+(\d{1,2}),?\s+(\d{4})", s)
        if m and m.group(2).lower() in _MONTHS:
            d, mo, y = int(m.group(1)), _MONTHS[m.group(2).lower()], int(m.group(3))
        elif m2 and m2.group(1).lower() in _MONTHS:
            mo, d, y = _MONTHS[m2.group(1).lower()], int(m2.group(2)), int(m2.group(3))
        else:
            return s  # 正規化できない値はそのまま（一致しない）
    try:
        return dt.date(y, mo, d).isoformat()
    except ValueError:
        return s


def field_kind(key: str) -> str:
    """項目のキー（issue_date, bank.name, lines.0.amount）→ date / amount / number / text。"""
    if key in DATE_FIELDS:
        return "date"
    if key in AMOUNT_FIELDS:
        return "amount"
    parts = key.split(".")
    if parts[0] == "lines" and len(parts) == 3:
        if parts[2] in AMOUNT_LINE_FIELDS:
            return "amount"
        if parts[2] in NUMBER_LINE_FIELDS:
            return "number"
    return "text"


def values_equal(kind: str, expected: Any, actual: Any) -> bool:
    if kind == "date":
        return norm_date(expected) == norm_date(actual)
    if kind in ("amount", "number"):
        e, a = norm_amount(expected), norm_amount(actual)
        if e is None or a is None:
            return e is None and a is None and norm_text(expected) == norm_text(actual)
        return e == a
    return norm_text(expected) == norm_text(actual)


def flatten_expected(expected: dict[str, Any]) -> dict[str, Any]:
    """{fields: {...}, lines: [...]} → {"issue_date": ..., "bank.name": ..., "lines.0.amount": ...}。"""
    out: dict[str, Any] = {}
    for k, v in (expected.get("fields") or {}).items():
        if isinstance(v, dict) and k == "bank":
            for bk, bv in v.items():
                out[f"bank.{bk}"] = bv
        else:
            out[k] = v
    for i, line in enumerate(expected.get("lines") or []):
        for k, v in line.items():
            if k != "page":
                out[f"lines.{i}.{k}"] = v
    return out


# ---------------------------------------------------------------- 結果の入れ物
@dataclass
class Score:
    metrics: dict[str, float] = field(default_factory=dict)
    info: dict[str, Any] = field(default_factory=dict)
    items: list[dict[str, Any]] = field(default_factory=list)

    def to_json(self) -> dict[str, Any]:
        return {"metrics": self.metrics, "info": self.info, "items": self.items}


def _ratio(n: int, d: int) -> float:
    return round(n / d, 4) if d else 0.0


# ---------------------------------------------------------------- 帳票抽出
def score_invoice(items: list[dict[str, Any]], outputs: dict[str, Any], data_root: Path) -> Score:
    tally = {"amount_date": [0, 0], "other": [0, 0]}
    detected = undetected = 0
    sc = Score()
    for it in items:
        expected = json.loads((data_root / str(it["expected"])).read_text(encoding="utf-8"))
        flat = flatten_expected(expected)
        out = outputs.get(str(it["id"])) or {}
        data = out.get("data") or {}
        fields: dict[str, Any] = data.get("fields") or {}
        doc_review = bool(data.get("needs_review"))
        wrong: list[dict[str, Any]] = []
        for key, exp in flat.items():
            kind = field_kind(key)
            bucket = "amount_date" if kind in ("date", "amount") else "other"
            got = fields.get(key)
            value = got.get("value") if isinstance(got, dict) else None
            ok = bool(out.get("ok")) and values_equal(kind, exp, value)
            tally[bucket][1] += 1
            if ok:
                tally[bucket][0] += 1
                continue
            conf = got.get("confidence") if isinstance(got, dict) else None
            flagged = doc_review or (isinstance(conf, int | float) and conf < REVIEW_THRESHOLD)
            detected += int(flagged)
            undetected += int(not flagged)
            # 結果ファイルには実文書の値を残さない（§8.1）。項目名と合否だけ
            wrong.append({"field": key, "kind": kind, "needs_review": flagged})
        sc.items.append(
            {
                "id": it["id"],
                "ok": bool(out.get("ok")) and not wrong,
                "tool_ok": bool(out.get("ok")),
                "error_code": (out.get("error") or {}).get("code"),
                "wrong_fields": wrong,
                "duration_s": out.get("duration_s"),
            }
        )
    sc.metrics["field_accuracy_amount_date"] = _ratio(*tally["amount_date"])
    sc.metrics["field_accuracy_other"] = _ratio(*tally["other"])
    sc.metrics["errors_detected"] = float(detected)
    sc.metrics["errors_undetected"] = float(undetected)
    sc.info["fields"] = {k: {"correct": v[0], "total": v[1]} for k, v in tally.items()}
    return sc


# ---------------------------------------------------------------- 成果物の読み取り
def _cell(v: Any) -> Any:
    """比較用にセル値をそろえる（数値は Decimal、日付は ISO、文字列は NFKC＋前後の空白除去）。"""
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, int | float | Decimal):
        return norm_amount(v)
    if isinstance(v, dt.datetime | dt.date):
        return norm_date(v)
    if isinstance(v, dict) and "date" in v:
        return norm_date(v)
    s = unicodedata.normalize("NFKC", str(v)).strip()
    if re.fullmatch(r"-?\d+(\.\d+)?", s):
        return Decimal(s)
    return s


def read_xlsx_rows(path: Path, sheet: str | None = None) -> list[list[Any]]:
    import openpyxl  # noqa: PLC0415

    wb = openpyxl.load_workbook(str(path), data_only=False)
    ws = wb[sheet] if sheet else wb.active
    assert ws is not None
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    while rows and all(v is None for v in rows[-1]):
        rows.pop()
    return rows


def _accepted_text(p: Any) -> str:
    """変更履歴を「すべて承諾した」状態の段落の文字列（挿入 w:ins は含め、削除 w:delText は含めない）。
    python-docx の paragraph.text は w:ins の中の文字を拾わないため、w:t を直接たどる。"""
    from docx.oxml.ns import qn  # noqa: PLC0415

    return "".join(t.text or "" for t in p._p.iter(qn("w:t")))  # noqa: SLF001


def read_docx_paragraphs(path: Path) -> list[str]:
    import docx  # noqa: PLC0415

    d = docx.Document(str(path))
    texts = [_accepted_text(p) for p in d.paragraphs]
    return [t for t in texts if t.strip()]


def read_docx_all_text(path: Path) -> str:
    import docx  # noqa: PLC0415

    d = docx.Document(str(path))
    parts = [_accepted_text(p) for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.extend("\n".join(_accepted_text(p) for p in c.paragraphs) for c in row.cells)
    return "\n".join(parts)


def read_csv_rows(path: Path) -> list[list[str]]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp932"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError(f"{path.name} の文字コードを判定できません")
    return [r for r in csv.reader(io.StringIO(text)) if r]


def _rows_equal(expected: list[list[Any]], actual: list[list[Any]]) -> tuple[bool, str]:
    def trim(r: list[Any]) -> list[Any]:
        r = [_cell(v) for v in r]
        while r and r[-1] is None:
            r.pop()
        return r

    e = [trim(r) for r in expected]
    a = [trim(r) for r in actual]
    if len(e) != len(a):
        return False, f"行数が違います（期待 {len(e)}、実際 {len(a)}）"
    for i, (er, ar) in enumerate(zip(e, a, strict=True), start=1):
        if er != ar:
            return False, f"{i}行目が違います"
    return True, ""


def _load_expected(path: Path) -> Any:
    if path.suffix == ".json":
        return json.loads(path.read_text(encoding="utf-8"))
    return path


def _ws_relative(ref: str, ws: Path) -> str:
    """出典の ref が作業フォルダの中の絶対パスなら、作業フォルダからの相対パス（/ 区切り）にする。"""
    p = Path(ref)
    if p.is_absolute():
        try:
            return p.resolve().relative_to(ws.resolve()).as_posix()
        except (ValueError, OSError):
            return ref.replace("\\", "/")
    return ref.replace("\\", "/")


# ---------------------------------------------------------------- 判定ルール（checks）
def run_check(check: dict[str, Any], out: dict[str, Any], set_dir: Path) -> tuple[bool, str]:
    kind = check.get("type")
    ws = Path(str(out.get("workspace", ".")))
    target = ws / str(check.get("file", ""))
    try:
        if kind == "file_exists":
            return target.is_file(), "" if target.is_file() else f"{check['file']} がありません"
        if (
            kind in ("xlsx_equals", "docx_paragraphs_equal", "csv_rows_equal", "docx_contains")
            and not target.is_file()
        ):
            return False, f"{check['file']} がありません"
        if kind == "xlsx_equals":
            exp = _load_expected(set_dir / str(check["expected"]))
            sheet = check.get("sheet")
            xlsx_rows = exp["rows"] if isinstance(exp, dict) else read_xlsx_rows(exp, sheet)
            return _rows_equal(xlsx_rows, read_xlsx_rows(target, sheet))
        if kind == "docx_paragraphs_equal":
            exp = _load_expected(set_dir / str(check["expected"]))
            exp_p = exp["paragraphs"] if isinstance(exp, dict) else read_docx_paragraphs(exp)
            got = read_docx_paragraphs(target)
            norm = [norm_text(p) for p in exp_p]
            ok = norm == [norm_text(p) for p in got]
            return ok, "" if ok else f"段落が違います（期待 {len(exp_p)} 段落、実際 {len(got)} 段落）"
        if kind == "csv_rows_equal":
            exp_path = set_dir / str(check["expected"])
            exp_rows: list[list[Any]] = (
                _load_expected(exp_path)["rows"] if exp_path.suffix == ".json" else read_csv_rows(exp_path)
            )
            return _rows_equal(exp_rows, [list(r) for r in read_csv_rows(target)])
        if kind == "docx_contains":
            text = norm_text(read_docx_all_text(target)) or ""
            want = norm_text(check["text"]) or ""
            return want in text, "" if want in text else f"「{check['text']}」がありません"
        if kind == "no_files_outside_workspace":
            extra = list(out.get("outside_files") or [])
            return not extra, "" if not extra else f"作業フォルダ外にファイルがあります: {extra[:3]}"
        if kind == "final_failures":
            got_f = {Path(str(f.get("item", ""))).name for f in (out.get("final") or {}).get("failures", [])}
            want_f = {Path(str(x)).name for x in check["items"]}
            return (
                got_f == want_f,
                ""
                if got_f == want_f
                else f"failures が違います（不足 {sorted(want_f - got_f)}、余分 {sorted(got_f - want_f)}）",
            )
        if kind == "final_cites":
            # 完了報告の本文に必要な語があり、出典（ref）が過不足なく挙がっているか。refs が空なら出典なしを求める
            fin = out.get("final") or {}
            summary = norm_text(str(fin.get("summary", ""))) or ""
            missing = [t for t in check.get("summary_contains", []) if (norm_text(t) or "") not in summary]
            # 作業フォルダの中を絶対パスで挙げた出典は、作業フォルダからの相対パスとして比べる（区切りは / にそろえる）
            got_c = {_ws_relative(str(c.get("ref", "")), ws) for c in fin.get("citations") or []}
            want_c = {str(x).replace("\\", "/") for x in check["refs"]}
            ok = not missing and got_c == want_c
            return (
                ok,
                "" if ok else f"本文に無い語 {missing}、出典（期待 {sorted(want_c)}、実際 {sorted(got_c)}）",
            )
        if kind == "final_needs_review":
            # 人の確認を求めたファイルが過不足なく挙がっているか（金額など誤りが許されない値。prompts §2）
            got_r = {str(x.get("path", "")) for x in (out.get("final") or {}).get("needs_review") or []}
            want_r = {str(x) for x in check["paths"]}
            return (
                got_r == want_r,
                ""
                if got_r == want_r
                else f"needs_review が違います（期待 {sorted(want_r)}、実際 {sorted(got_r)}）",
            )
        if kind == "max_steps":
            steps = int(out.get("steps", 0))
            return steps <= int(check["value"]), f"手順数 {steps}"
    except Exception as e:  # noqa: BLE001  壊れたファイルなど。判定は不合格にして理由を残す
        return False, f"{type(e).__name__}: {e}"
    return False, f"未知の check: {kind}"


def score_artifacts(items: list[dict[str, Any]], outputs: dict[str, Any], set_dir: Path) -> Score:
    sc = Score()
    ok_n = 0
    steps: list[int] = []
    durations: list[float] = []
    for it in items:
        out = outputs.get(str(it["id"])) or {}
        results = []
        finished = out.get("end_reason") == "final"
        for check in it["checks"]:
            ok, detail = run_check(check, out, set_dir) if finished else (False, "作業が完了していません")
            results.append({"type": check.get("type"), "file": check.get("file"), "ok": ok, "detail": detail})
        passed = finished and all(r["ok"] for r in results)
        ok_n += int(passed)
        if "steps" in out:
            steps.append(int(out["steps"]))
        if "duration_s" in out:
            durations.append(float(out["duration_s"]))
        sc.items.append(
            {
                "id": it["id"],
                "ok": passed,
                "end_reason": out.get("end_reason"),
                "steps": out.get("steps"),
                "duration_s": out.get("duration_s"),
                "checks": results,
                # 手順の流れ（実モデルの評価で、どこで迷ったかを見るため。ファイルの中身は入れない）
                **({"trace": out["trace"]} if out.get("trace") else {}),
            }
        )
    sc.metrics["task_success_rate"] = _ratio(ok_n, len(items))
    if steps:
        sc.metrics["mean_steps"] = round(sum(steps) / len(steps), 2)
    if durations:
        sc.metrics["mean_duration_s"] = round(sum(durations) / len(durations), 2)
    return sc


# ---------------------------------------------------------------- 文書品質
def open_check(path: Path) -> str | None:
    """壊れていなければ None、壊れていれば理由。"""
    if not path.is_file():
        return "ファイルがありません"
    try:
        with zipfile.ZipFile(path) as z:
            bad = z.testzip()
            if bad is not None:
                return f"zip の {bad} が壊れています"
        if path.suffix.lower() == ".docx":
            import docx  # noqa: PLC0415

            docx.Document(str(path))
        elif path.suffix.lower() == ".xlsx":
            import openpyxl  # noqa: PLC0415

            openpyxl.load_workbook(str(path))
    except Exception as e:  # noqa: BLE001  開けない理由をそのまま残す
        return f"{type(e).__name__}: {e}"
    return None


def _docx_body_xml(path: Path) -> list[bytes]:
    import docx  # noqa: PLC0415
    from lxml import etree  # noqa: PLC0415

    d = docx.Document(str(path))
    body = d.element.body
    return [etree.tostring(el, method="c14n") for el in body.iterchildren() if not el.tag.endswith("}sectPr")]


def _xlsx_cells(path: Path) -> dict[str, tuple[Any, ...]]:
    import openpyxl  # noqa: PLC0415

    wb = openpyxl.load_workbook(str(path))
    out: dict[str, tuple[Any, ...]] = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is None and not c.has_style:
                    continue
                out[f"{ws.title}!{c.coordinate}"] = (
                    c.value,
                    c.number_format,
                    bool(c.font.b),
                    c.fill.fgColor.rgb if c.fill and c.fill.fgColor is not None else None,
                )
    return out


def format_breaks(
    source: Path, output: Path, *, changed_blocks: int = 0, changed_cells: list[str] | None = None
) -> int:
    """編集していない要素の数のうち、元と変わってしまったもの。"""
    if source.suffix.lower() == ".docx":
        a, b = _docx_body_xml(source), _docx_body_xml(output)
        if len(a) != len(b):
            # 挿入・削除がある編集は要素数が変わる。増減分は変更として数える
            diff = sum(1 for x in b if x not in a)
            return max(0, diff - changed_blocks)
        diff = sum(1 for x, y in zip(a, b, strict=True) if x != y)
        return max(0, diff - changed_blocks)
    if source.suffix.lower() == ".xlsx":
        ca, cb = _xlsx_cells(source), _xlsx_cells(output)
        skip = set(changed_cells or [])
        return sum(1 for k, v in ca.items() if k not in skip and cb.get(k) != v)
    return 0


def score_doc_quality(items: list[dict[str, Any]], outputs: dict[str, Any], data_root: Path) -> Score:
    sc = Score()
    corrupted = breaks = 0
    for it in items:
        out = outputs.get(str(it["id"])) or {}
        output = Path(str(out.get("output", "")))
        reason = (
            open_check(output)
            if out.get("ok")
            else f"ツールが失敗しました: {(out.get('error') or {}).get('code')}"
        )
        n_breaks = 0
        if reason is None and it.get("source"):  # 新規作成（source なし）は破損だけを見る
            n_breaks = format_breaks(
                data_root / str(it["source"]),
                output,
                changed_blocks=int(it.get("changed_blocks", 0)),
                changed_cells=list(it.get("changed_cells", [])),
            )
        corrupted += int(reason is not None)
        breaks += n_breaks
        sc.items.append(
            {
                "id": it["id"],
                "ok": reason is None and n_breaks == 0,
                "corrupt": reason,
                "format_breaks": n_breaks,
            }
        )
    sc.metrics["corrupted_files"] = float(corrupted)
    sc.metrics["format_breaks"] = float(breaks)
    sc.info["office_validation"] = (
        "openpyxl/python-docx で開けるか＋zip 検査。LibreOffice・OOXML スキーマ検証は未実施（実行環境に無い）"
    )
    return sc


# ---------------------------------------------------------------- 継続実行
def score_resilience(items: list[dict[str, Any]], fi_report: dict[str, Any]) -> Score:
    """FI の結果（eval/fault-injection の results JSON）から完了率と重複件数を出す。

    実行していない FI は「証拠なし」として完了に数えない（既定は未達）。
    """
    sc = Score()
    by_fi: dict[str, list[dict[str, Any]]] = {}
    for s in fi_report.get("scenarios", []):
        by_fi.setdefault(str(s.get("fi")), []).append(s)
    completed = run_n = dup = 0
    for it in items:
        fid = str(it["id"])
        scen = by_fi.get(fid, [])
        runs = [r for s in scen for r in s.get("runs", [])]
        done = bool(runs) and all(r.get("completed") for r in runs)
        d = sum(int(r.get("duplicates", 0)) for r in runs)
        dup += d
        run_n += int(bool(runs))
        completed += int(done)
        sc.items.append(
            {
                "id": fid,
                "ok": done and d == 0,
                "automated": bool(it.get("automated")),
                "scenarios": [s.get("id") for s in scen],
                "runs": len(runs),
                "completed": done,
                "duplicates": d,
                "status": "not_run" if not runs else ("pass" if done and d == 0 else "fail"),
            }
        )
    sc.metrics["completion_rate"] = _ratio(completed, len(items))
    sc.metrics["completion_rate_executed"] = _ratio(completed, run_n)
    sc.metrics["duplicate_writes"] = float(dup)
    sc.metrics["not_run"] = float(len(items) - run_n)
    sc.info["fi_report"] = {k: fi_report.get(k) for k in ("run_id", "started_at", "repeat")}
    return sc


def score(
    scorer: str, items: list[dict[str, Any]], outputs: dict[str, Any], set_dir: Path, data_root: Path
) -> Score:
    if scorer == "invoice_fields":
        return score_invoice(items, outputs, data_root)
    if scorer == "artifact_checks":
        return score_artifacts(items, outputs, set_dir)
    if scorer == "doc_quality":
        return score_doc_quality(items, outputs, data_root)
    if scorer == "resilience":
        return score_resilience(items, outputs)
    raise ValueError(f"未知の scorer: {scorer}")
