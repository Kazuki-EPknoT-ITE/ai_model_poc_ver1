"""合成データの生成（local_llm_test_spec.md §8.1。架空の会社名・船名・口座だけを使う）。

評価セットの data_root: synthetic と、障害注入試験（eval/fault-injection）の両方から使う。
data_hash で照合できるよう、同じ引数なら同じバイト列になるようにする（PDF は invariant、
Office 文書は zip の時刻と docProps の日時を固定する）。
"""

from __future__ import annotations

import json
import re
import shutil
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import Any

FIXED_ZIP_TIME = (2026, 1, 1, 0, 0, 0)
FIXED_ISO = "2026-01-01T00:00:00Z"


# ---------------------------------------------------------------- 請求書（PDF）
def invoice_truth(n: int) -> dict[str, Any]:
    """n 番目の合成請求書の正解値（items の expected と同じ形）。"""
    amount = n * 1000
    return {
        "fields": {
            "issue_date": f"2026-09-{(n - 1) % 28 + 1:02d}",
            "due_date": "2026-10-31",
            "issuer_name": f"Example Shipping Co. {n}",
            "bill_to": "Sample Marine K.K.",
            "vessel_name": "MV EXAMPLE STAR",
            "currency": "USD",
            "subtotal": str(amount),
            "tax": "0",
            "total": str(amount),
            "bank.name": "Example Bank",
            "bank.account_no": f"{1000000 + n}",
        },
        "lines": [
            {"description": "Charter hire", "qty": "1", "unit_price": str(amount), "amount": str(amount)}
        ],
    }


def invoice_pdf(path: Path, n: int, *, encrypt_password: str | None = None) -> None:
    from reportlab.lib.pagesizes import A4  # type: ignore[import-untyped]  # noqa: PLC0415
    from reportlab.pdfgen import canvas  # type: ignore[import-untyped]  # noqa: PLC0415

    t = invoice_truth(n)["fields"]
    amount = f"{n * 1000:,.2f}"
    kw: dict[str, Any] = {"pagesize": A4, "invariant": 1}
    if encrypt_password is not None:
        from reportlab.lib.pdfencrypt import StandardEncryption  # type: ignore[import-untyped]  # noqa: PLC0415

        kw["encrypt"] = StandardEncryption(encrypt_password, canPrint=0)
    c = canvas.Canvas(str(path), **kw)
    _, h = A4
    c.setFont("Helvetica", 11)
    lines = [
        "INVOICE",
        f"Issuer: {t['issuer_name']}",
        f"Bill To: {t['bill_to']}",
        f"Vessel: {t['vessel_name']}",
        f"Issue Date: {t['issue_date']}   Due Date: {t['due_date']}",
        f"Currency: {t['currency']}",
        f"Charter hire   1 x {amount} = {amount}",
        f"Subtotal: {amount}",
        "Tax: 0.00",
        f"Total: {amount}",
        f"Bank: {t['bank.name']}  Account No: {t['bank.account_no']}",
    ]
    y = h - 72
    for line in lines:
        c.drawString(72, y, line)
        y -= 18
    c.showPage()
    c.save()


def make_invoices(
    out: Path,
    count: int,
    *,
    broken: dict[str, list[int]] | None = None,
    expected: bool = True,
    start: int = 1,
) -> list[str]:
    """INV-NNN.pdf を count 件作る。broken で壊れたファイルを混ぜる（FI-18、long_batch）。

    broken: {"header": [番号...], "empty": [...], "encrypted": [...]}（番号はファイル名の NNN）。戻り値は壊したファイル名。
    start: 中身（請求元・金額など）の通し番号の始まり。ファイル名は 1 からのまま。学習データ（ml/datagen）は
    評価セットと同じファイルにならないよう、大きな番号から作る（test_spec §3.1 の重複照合）。
    """
    out.mkdir(parents=True, exist_ok=True)
    bad: list[str] = []
    kinds = {n: k for k, ns in (broken or {}).items() for n in ns}
    width = max(3, len(str(count)))
    for n in range(1, count + 1):
        name = f"INV-{n:0{width}d}.pdf"
        p = out / name
        kind = kinds.get(n)
        if kind == "empty":
            p.write_bytes(b"")
        elif kind == "encrypted":
            invoice_pdf(p, n + start - 1, encrypt_password="secret-pw")  # noqa: S106  合成データのパスワード
        else:
            invoice_pdf(p, n + start - 1)
            if kind == "header":
                data = bytearray(p.read_bytes())
                data[0:64] = b"\x00GARBAGE-HEADER" + b"\xff" * 49
                p.write_bytes(bytes(data))
        if kind:
            bad.append(name)
        elif expected:
            (out / "expected").mkdir(exist_ok=True)
            (out / "expected" / f"{p.stem}.json").write_text(
                json.dumps(invoice_truth(n + start - 1), ensure_ascii=False, indent=1), encoding="utf-8"
            )
    return bad


def invoice_model_output(n: int, *, overrides: dict[str, Any] | None = None) -> dict[str, Any]:
    """モック Gateway が invoice_extract に返す「モデルの抽出結果」（invoice.v1 の response_schema の形）。"""
    t = invoice_truth(n)["fields"]
    amount = float(n * 1000)

    def f(v: Any) -> dict[str, Any]:
        return {"value": v, "page": 1 if v is not None else None}

    out: dict[str, Any] = {
        "issue_date": f(t["issue_date"]),
        "due_date": f(t["due_date"]),
        "issuer_name": f(t["issuer_name"]),
        "issuer_reg_no": f(None),
        "bill_to": f(t["bill_to"]),
        "vessel_name": f(t["vessel_name"]),
        "voyage_no": f(None),
        "currency": f(t["currency"]),
        "subtotal": f(amount),
        "tax": f(0),
        "total": f(amount),
        "bank": {
            "name": f(t["bank.name"]),
            "branch": f(None),
            "account_type": f(None),
            "account_no": f(t["bank.account_no"]),
            "swift": f(None),
        },
        "lines": [
            {"description": "Charter hire", "qty": 1, "unit_price": amount, "amount": amount, "page": 1}
        ],
    }
    for key, value in (overrides or {}).items():
        if key.startswith("bank."):
            out["bank"][key[5:]] = f(value)
        else:
            out[key] = f(value)
    return out


# ---------------------------------------------------------------- Office 文書（決定的に保存する）
def stabilize_zip(path: Path) -> None:
    """zip の時刻・順序と docProps/core.xml の日時を固定し、同じ内容なら同じバイト列にする。"""
    with zipfile.ZipFile(path) as z:
        entries = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as z:
        for name, data in entries:
            if name == "docProps/core.xml":
                text = data.decode("utf-8")
                text = re.sub(
                    r"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:)", rf"\g<1>{FIXED_ISO}\g<2>", text
                )
                data = text.encode("utf-8")
            info = zipfile.ZipInfo(name, date_time=FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, data)
    tmp.replace(path)


def make_report_docx(path: Path) -> None:
    import docx  # noqa: PLC0415

    d = docx.Document()
    d.add_heading("入渠報告書（合成データ）", level=1)
    d.add_paragraph("船名: MV EXAMPLE STAR")
    d.add_paragraph("報告日: 2026-08-31")
    d.add_heading("費用", level=2)
    t = d.add_table(rows=3, cols=2)
    for r, (k, v) in enumerate((("項目", "金額"), ("入渠費", "1,000,000"), ("塗装費", "500,000"))):
        t.cell(r, 0).text = k
        t.cell(r, 1).text = v
    p = d.add_paragraph("以上、報告します。")
    p.runs[0].bold = True
    d.save(str(path))
    stabilize_zip(path)


def make_budget_xlsx(path: Path) -> None:
    import openpyxl  # noqa: PLC0415
    from openpyxl.styles import Font, PatternFill  # noqa: PLC0415

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "予算"
    ws.append(["項目", "金額", "備考"])
    for row in (["入渠費", 1000000, "合成"], ["塗装費", 500000, "合成"], ["検査費", 250000, "合成"]):
        ws.append(row)
    ws["A5"] = "合計"
    ws["B5"] = "=SUM(B2:B4)"
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDDDDD")
    for r in range(2, 6):
        ws.cell(r, 2).number_format = "#,##0"
    wb.save(str(path))
    stabilize_zip(path)


# ---------------------------------------------------------------- セットごとの生成
def _gen_invoices(spec: dict[str, Any], root: Path) -> None:
    make_invoices(root, int(spec.get("count", 3)), broken=spec.get("broken"))


def _gen_tool_tasks(spec: dict[str, Any], root: Path) -> None:
    ws = root / "workspaces"
    for task in spec.get("tasks", []):
        tid = str(task["id"])
        d = ws / tid
        d.mkdir(parents=True, exist_ok=True)
        kind = task["kind"]
        if kind == "invoices":
            make_invoices(d / "請求書", int(task.get("count", 3)), expected=False)
        elif kind == "memo":
            (d / "メモ.txt").write_text("品名,数量\nロープ,4\nペンキ,12\n", encoding="utf-8")
        elif kind == "empty":
            (d / ".keep").write_text("", encoding="utf-8")
        else:
            raise ValueError(f"未知の tasks.kind: {kind}")


def _gen_doc_sources(spec: dict[str, Any], root: Path) -> None:
    src = root / "sources"
    src.mkdir(parents=True, exist_ok=True)
    make_report_docx(src / "report.docx")
    make_budget_xlsx(src / "budget.xlsx")


GENERATORS: dict[str, Callable[[dict[str, Any], Path], None]] = {
    "invoices": _gen_invoices,
    "tool_tasks": _gen_tool_tasks,
    "doc_sources": _gen_doc_sources,
}


def generate(spec: dict[str, Any], root: Path) -> None:
    kind = str(spec.get("kind"))
    if kind not in GENERATORS:
        raise ValueError(f"未知の generator.kind: {kind}")
    if root.exists():
        shutil.rmtree(root)
    root.mkdir(parents=True)
    GENERATORS[kind](spec, root)
