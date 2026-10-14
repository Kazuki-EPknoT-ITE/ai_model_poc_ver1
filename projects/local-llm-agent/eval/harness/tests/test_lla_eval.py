"""評価ハーネスの Python 側（manifest の検証・合成データ・採点・合否表・台帳向け出力）。"""

from __future__ import annotations

import json
import shutil
from pathlib import Path
from typing import Any

import pytest
import yaml
from lla_eval.__main__ import main
from lla_eval.manifest import ManifestError, Metric, load_set, parse_manifest
from lla_eval.report import (
    VERDICT_FAIL,
    VERDICT_NO_EVIDENCE,
    VERDICT_PASS,
    build_result,
    registry_rows,
    table,
)
from lla_eval.scoring import norm_amount, norm_date, run_check, score
from lla_eval.synth import invoice_model_output, invoice_truth

SETS = Path(__file__).resolve().parents[2] / "sets"


def _manifest(**over: Any) -> dict[str, Any]:
    raw = yaml.safe_load((SETS / "invoice_extraction_sample" / "manifest.yaml").read_text(encoding="utf-8"))
    raw.update(over)
    return dict(raw)


def test_metric_targets() -> None:
    assert Metric("a", ">= 0.98").check(0.98) is True
    assert Metric("a", ">= 0.98").check(0.9799) is False
    assert Metric("a", "== 0").check(0) is True
    assert Metric("a", "< 5").check(None) is None  # 証拠なし


def test_manifest_problems_are_collected() -> None:
    with pytest.raises(ManifestError) as e:
        parse_manifest({**_manifest(), "unknown_key": 1, "version": "1", "runner": "nope"})
    joined = " / ".join(e.value.problems)
    assert "未知の項目 unknown_key" in joined
    assert "version" in joined
    assert len(e.value.problems) >= 3
    with pytest.raises(ManifestError):
        parse_manifest(["not", "a", "mapping"])


def test_normalizers() -> None:
    assert norm_amount("¥1,000") == norm_amount("1000.00")
    assert norm_date("2026/9/3") == norm_date("2026-09-03")


def _copy_set(tmp_path: Path, name: str) -> Path:
    dst = tmp_path / name
    shutil.copytree(SETS / name, dst)
    shutil.copytree(SETS / "_train", tmp_path / "_train")
    return dst


def test_data_hash_mismatch_refuses_to_run(tmp_path: Path) -> None:
    set_dir = _copy_set(tmp_path, "invoice_extraction_sample")
    load_set(set_dir, tmp_path / "cache")  # 登録済みのハッシュと一致する
    raw = yaml.safe_load((set_dir / "manifest.yaml").read_text(encoding="utf-8"))
    raw["data_hash"] = "sha256:" + "0" * 64
    (set_dir / "manifest.yaml").write_text(yaml.safe_dump(raw, allow_unicode=True), encoding="utf-8")
    with pytest.raises(ManifestError) as e:
        load_set(set_dir, tmp_path / "cache")
    assert "data_hash" in " ".join(e.value.problems)


def test_train_overlap_refuses_to_run(tmp_path: Path) -> None:
    set_dir = _copy_set(tmp_path, "invoice_extraction_sample")
    loaded = load_set(set_dir, tmp_path / "cache")
    # 評価データのファイルを学習データの一覧に入れる → 実行しない
    from lla_eval.manifest import data_files, file_sha256

    first = data_files(loaded.data_root)[0]
    index = tmp_path / "_train" / "train_hashes.sample.txt"
    index.write_text(index.read_text(encoding="utf-8") + file_sha256(first) + "\n", encoding="utf-8")
    with pytest.raises(ManifestError) as e:
        load_set(set_dir, tmp_path / "cache")
    assert "学習データと重複" in e.value.problems[0]


def test_empty_files_are_not_overlap_evidence(tmp_path: Path) -> None:
    from lla_eval.manifest import file_sha256, overlap_files

    (tmp_path / "a.pdf").write_bytes(b"")
    (tmp_path / "b.pdf").write_bytes(b"%PDF-1.4")
    # 空のファイル（壊れたファイルの例）はどこでも同じハッシュになるため、照合の対象にしない
    assert [p.name for p in overlap_files(tmp_path)] == ["b.pdf"]
    set_dir = _copy_set(tmp_path / "s", "invoice_extraction_sample")
    loaded = load_set(set_dir, tmp_path / "cache")
    (loaded.data_root / "empty.pdf").write_bytes(b"")
    index = tmp_path / "s" / "_train" / "train_hashes.sample.txt"
    index.write_text(
        index.read_text(encoding="utf-8") + file_sha256(tmp_path / "a.pdf") + "\n", encoding="utf-8"
    )
    from lla_eval.manifest import split_check

    assert split_check(loaded.manifest, set_dir, loaded.data_root)["status"] == "ok"


def _invoice_outputs(n_items: int, *, wrong: dict[int, dict[str, Any]]) -> dict[str, Any]:
    outputs: dict[str, Any] = {}
    for n in range(1, n_items + 1):
        model = invoice_model_output(n, overrides=wrong.get(n))
        fields: dict[str, Any] = {}
        for k, v in model.items():
            if k == "bank":
                for bk, bv in v.items():
                    fields[f"bank.{bk}"] = bv
            elif k == "lines":
                for i, line in enumerate(v):
                    for lk, lv in line.items():
                        if lk != "page":
                            fields[f"lines.{i}.{lk}"] = {"value": lv}
            elif isinstance(v, dict):
                fields[k] = v
        outputs[f"INV-{n:03d}"] = {"ok": True, "data": {"fields": fields, "needs_review": n in wrong}}
    return outputs


def test_invoice_scoring_and_report(tmp_path: Path) -> None:
    set_dir = _copy_set(tmp_path, "invoice_extraction_sample")
    loaded = load_set(set_dir, tmp_path / "cache")
    m = loaded.manifest
    perfect = score(m.scorer, loaded.items, _invoice_outputs(3, wrong={}), set_dir, loaded.data_root)
    assert perfect.metrics["field_accuracy_amount_date"] == 1.0
    assert perfect.metrics["field_accuracy_other"] == 1.0
    result = build_result(loaded.to_json(), m, {"run_id": "r1", "model_id": "mock"}, perfect)
    assert result["verdict"] == VERDICT_PASS
    assert registry_rows(result)[0] == {
        "eval_set": "invoice_extraction_sample",
        "metric": "field_accuracy_amount_date",
        "value": 1.0,
    }

    # 2枚目の合計を誤らせる（needs_review が立っているので「検出できた誤り」）
    sc = score(
        m.scorer, loaded.items, _invoice_outputs(3, wrong={2: {"total": 1.0}}), set_dir, loaded.data_root
    )
    assert sc.metrics["field_accuracy_amount_date"] < 1.0
    assert sc.metrics["errors_detected"] == 1.0 and sc.metrics["errors_undetected"] == 0.0
    bad = build_result(loaded.to_json(), m, {"run_id": "r2"}, sc)
    assert bad["verdict"] == VERDICT_FAIL
    # 結果ファイルに実文書の値は残さない（項目名だけ）
    dumped = json.dumps(bad, ensure_ascii=False)
    assert invoice_truth(2)["fields"]["issuer_name"] not in dumped
    assert "INV-002" in table(bad) and "未達" in table(bad)


def test_no_evidence_is_not_a_pass(tmp_path: Path) -> None:
    set_dir = _copy_set(tmp_path, "invoice_extraction_sample")
    loaded = load_set(set_dir, tmp_path / "cache")
    m = parse_manifest(_manifest(metrics=[{"name": "something_unmeasured", "target": ">= 1"}]))
    sc = score(m.scorer, loaded.items, {}, set_dir, loaded.data_root)
    result = build_result(loaded.to_json(), m, {"run_id": "r"}, sc)
    assert result["metrics"][0]["verdict"] == VERDICT_NO_EVIDENCE
    assert result["verdict"] == VERDICT_FAIL


def test_cli_finalize_writes_result_and_registry(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    set_dir = _copy_set(tmp_path, "invoice_extraction_sample")
    record = tmp_path / "record.json"
    record.write_text(
        json.dumps({"run_id": "cli-1", "model_id": "mock", "outputs": _invoice_outputs(3, wrong={})}),
        encoding="utf-8",
    )
    out = tmp_path / "out"
    rc = main(
        [
            "finalize",
            str(set_dir),
            "--record",
            str(record),
            "--out",
            str(out),
            "--cache",
            str(tmp_path / "c"),
            "--strict",
        ]
    )
    assert rc == 0
    last = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    assert last["verdict"] == VERDICT_PASS
    reg = json.loads((out / "cli-1.registry.json").read_text(encoding="utf-8"))
    # そのまま POST /admin/v1/models/{id}/evals の results に渡せる形
    assert {"eval_set", "metric", "value"} == set(reg[0])
    # 壊れた manifest は終了コード 2 と問題の一覧
    (set_dir / "manifest.yaml").write_text("name: x\n", encoding="utf-8")
    assert main(["load", str(set_dir), "--cache", str(tmp_path / "c")]) == 2
    assert "error" in json.loads(capsys.readouterr().out.strip().splitlines()[-1])


def test_final_cites_check(tmp_path: Path) -> None:
    fin = {
        "summary": "国内出張の日当は１日3,000円です。",
        "citations": [{"kind": "rag", "ref": "規程/旅費規程.docx"}],
    }
    out = {"workspace": str(tmp_path), "final": fin}
    ok = {"type": "final_cites", "summary_contains": ["日当は1日3,000円"], "refs": ["規程/旅費規程.docx"]}
    assert run_check(ok, out, tmp_path)[0]
    assert not run_check({**ok, "summary_contains": ["5,000円"]}, out, tmp_path)[0]
    assert not run_check({**ok, "refs": []}, out, tmp_path)[0]
    # 作業フォルダの中の絶対パスで挙げた出典も、相対パスの期待と一致する（外の絶対パスは一致しない）
    (tmp_path / "経費.csv").write_text("x", encoding="utf-8")
    abs_out = {
        "workspace": str(tmp_path),
        "final": {"summary": "s", "citations": [{"kind": "file", "ref": str(tmp_path / "経費.csv")}]},
    }
    assert run_check({"type": "final_cites", "refs": ["経費.csv"]}, abs_out, tmp_path)[0]
    other = {**abs_out, "workspace": str(tmp_path / "sub")}
    assert not run_check({"type": "final_cites", "refs": ["経費.csv"]}, other, tmp_path)[0]
    none = {"workspace": str(tmp_path), "final": {"summary": "見つかりませんでした。"}}
    assert run_check(
        {"type": "final_cites", "summary_contains": ["見つかりません"], "refs": []}, none, tmp_path
    )[0]


def test_final_needs_review_check(tmp_path: Path) -> None:
    fin = {"summary": "転記しました", "needs_review": [{"path": "送金明細.xlsx", "reason": "金額"}]}
    out = {"workspace": str(tmp_path), "final": fin}
    assert run_check({"type": "final_needs_review", "paths": ["送金明細.xlsx"]}, out, tmp_path)[0]
    assert not run_check({"type": "final_needs_review", "paths": []}, out, tmp_path)[0]
    assert run_check({"type": "final_needs_review", "paths": []}, {"final": {}}, tmp_path)[0]


def test_docx_reads_tracked_changes_as_accepted(tmp_path: Path) -> None:
    import docx
    from docx.oxml import parse_xml

    d = docx.Document()
    p = d.add_paragraph("支払期限は")
    ns = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    p._p.append(
        parse_xml(f'<w:del {ns} w:id="1" w:author="a"><w:r><w:delText>30日</w:delText></w:r></w:del>')
    )
    p._p.append(parse_xml(f'<w:ins {ns} w:id="2" w:author="a"><w:r><w:t>45日</w:t></w:r></w:ins>'))
    p.add_run("以内とする。")
    f = tmp_path / "a.docx"
    d.save(str(f))
    out = {"workspace": str(tmp_path)}
    assert run_check(
        {"type": "docx_contains", "file": "a.docx", "text": "支払期限は45日以内とする。"}, out, tmp_path
    )[0]
    assert not run_check({"type": "docx_contains", "file": "a.docx", "text": "30日"}, out, tmp_path)[0]
